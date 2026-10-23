"""수량첨부 예제 — 공정별·공종별 수량파일(설계내역 수량집계표 형식) 작성

  examples/수량첨부_예제/처리시설공사/{토공,가시설공,구조물공,포장공,부대공,관로공}.xlsx  (대죽 예제 수량)
  examples/수량첨부_예제/관공사/{토공,관로공,포장공}.xlsx · 배수설비 공사/배수설비.xlsx  (가상 예시 수량)
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)
from cpa.project import read_project  # noqa: E402

OUT = os.path.join(ROOT, "examples", "수량첨부_예제")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write(path, title, sections, two_row=False):
    """sections: [(제목, [(공종, 규격, 단위, 수량), ...])]"""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "수량집계표"
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    r = 3
    if two_row:   # 2단 머리글(병합) 형식
        for j, h in enumerate(["공 종", "규 격", "단 위", "수 량", "", "비 고"], 1):
            ws.cell(r, j, h)
        for j, h in enumerate(["", "", "", "설계", "계", ""], 1):
            ws.cell(r + 1, j, h)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        r += 2
    else:
        for j, h in enumerate(["공 종", "규 격", "단 위", "수 량", "비 고"], 1):
            c = ws.cell(r, j, h)
            c.font, c.fill, c.border = Font(bold=True), PatternFill("solid", fgColor="DDEBF7"), BOX
        r += 1
    for k, (head, rows) in enumerate(sections, 1):
        ws.cell(r, 1, f"{k}. {head}").font = Font(bold=True)
        r += 1
        for name, spec, unit, qty in rows:
            vals = [name, spec, unit, qty] + ([qty] if two_row else []) + [""]
            for j, v in enumerate(vals, 1):
                ws.cell(r, j, v).border = BOX
            r += 1
        ws.cell(r, 1, "소 계")
        r += 1
    for col, w in zip("ABCDEF", (30, 30, 8, 12, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.cell(r + 1, 1).alignment = Alignment()
    wb.save(path)


def main():
    p = read_project(os.path.join(ROOT, "examples", "예제_대죽공공폐수_입력.xlsx"))
    by = {}
    for it in p.items:
        by.setdefault(it.act, []).append((it.name, it.spec, it.unit, it.qty))
    d = os.path.join(OUT, "처리시설공사")
    write(os.path.join(d, "토공.xlsx"), "토공 수량집계표", [("터파기", by["C30"]), ("되메우기", by["C70"])])
    write(os.path.join(d, "가시설공.xlsx"), "가시설공 수량집계표", [("가시설 설치", by["C20"]), ("가시설 해체", by.get("C65", []))])
    write(os.path.join(d, "구조물공.xlsx"), "구조물공 수량집계표", [("TOC저감시설·밸브실", by["C40"])], two_row=True)
    write(os.path.join(d, "포장공.xlsx"), "포장공 수량집계표", [("포장 깨기", by["C10"]), ("포장 복구", by["C80"])])
    write(os.path.join(d, "부대공.xlsx"), "부대공 수량집계표", [("지장물", by["C25"]), ("기타", by["C90"])])
    write(os.path.join(d, "관로공.xlsx"), "구내배관 수량집계표", [("구내배관", by["C50"])])

    d = os.path.join(OUT, "관공사")
    write(os.path.join(d, "토공.xlsx"), "관로 토공 수량집계표", [
        ("터파기", [("터파기", "토사,육상,기계 0.7㎥", "㎥", 4200), ("터파기", "풍화암,육상,브레이커", "㎥", 650)]),
        ("되메우기", [("모래기초", "관기초", "㎥", 820), ("되메우기", "토사,다짐", "㎥", 3100)]),
        ("운반", [("사토운반", "L=10km", "㎥", 1750)])])
    write(os.path.join(d, "관로공.xlsx"), "관로공 수량집계표", [
        ("관 부설", [("고강성PVC이중벽관 부설", "D300mm", "m", 2350), ("고강성PVC이중벽관 부설", "D400mm", "m", 880)]),
        ("맨홀", [("소형맨홀 설치", "D900", "개소", 74)])])
    write(os.path.join(d, "포장공.xlsx"), "포장공 수량집계표", [
        ("포장 깨기", [("아스팔트포장 절단", "t=20cm 이하", "m", 6460), ("아스팔트포장 깨기", "t=20cm", "㎥", 520)]),
        ("포장 복구", [("보조기층", "t=20cm", "㎡", 2600), ("아스팔트 기층", "t=7cm", "㎡", 2600),
                    ("아스팔트 표층", "t=5cm", "㎡", 3900)])])
    write(os.path.join(OUT, "배수설비 공사", "배수설비.xlsx"), "배수설비 수량집계표", [
        ("배수설비", [("터파기", "토사,인력", "㎥", 380), ("PVC관 부설", "D150mm", "m", 1450),
                   ("오수받이 설치", "D400", "개소", 210), ("되메우기", "인력", "㎥", 330)])])
    print("작성:", OUT)


if __name__ == "__main__":
    main()
