"""입력서식(Excel) 정의·작성·읽기.

입력서식 시트
  0_사용안내 / 1_공사개요 / 2_비작업조건 / 3_수량 / 4_공정 / 5_유사사례
"""
import datetime as dt
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import refdata as R
from .schedule import Activity, parse_preds
from .workdays import Factors, Item

# ── 1_공사개요 항목: (키, 기본값, 설명, 선택목록) ─────────────────────────────
OVERVIEW = [
    ("■ 사업 정보", None, None, None),
    ("공사명", "", "보고서 표지·머리글에 사용", None),
    ("발주청", "", "", None),
    ("사업위치", "", "", None),
    ("시설유형(부록5)", "기타건축물", "부록5 산정공식 유형 (해당없음 가능)", list(R.FORMULAS) + ["해당없음"]),
    ("총공사비(백만원)", 0, "추정금액(추정가격+부가세+관급자재비)", None),
    ("착공예정일", dt.date(2026, 5, 1), "YYYY-MM-DD", None),
    ("■ 준비·정리기간 (비작업일수 미계상)", None, None, None),
    ("준비기간 공사유형", "상수도공사", "가이드라인 준비기간 예시 유형", list(R.PREP_PERIODS) + ["직접입력"]),
    ("준비기간(일)", None, "비우면 공사유형 예시값 적용", None),
    ("준비기간 산정사유", "", "보고서 문구", None),
    ("정리기간(일)", 30, "준공 전 1개월 범위", None),
    ("■ 비작업일수 조건", None, None, None),
    ("기상관측지점", "서울", "부록3 95개 관측지점 중 선택(최인접 행정지역)", sorted(R.WEATHER["stations"])),
    ("비작업일수 정수반올림", "Y", "월별 비작업일수 소수점 첫째자리 반올림", ["Y", "N"]),
    ("주40시간 하한(일/월)", 8, "가이드라인: 월 8일", None),
    ("동절기 공사중지", "추가검토", "미반영 / 반영 / 추가검토(미반영 기준 산정 + 반영 시 별도 검토)", ["미반영", "반영", "추가검토"]),
    ("동절기 중지월", "12,1,2", "쉼표 구분", None),
    ("준비·정리기간 동절기 중지", "Y", "Y: 동절기에는 준비·정리기간도 진행하지 않음 / N: 동절기 중에도 준비·정리 진행", ["Y", "N"]),
    ("■ 작업일수 보정 (부록4 1일작업량에 곱하는 계수)", None, None, None),
    ("현장여건", "보통", "", list(R.ADJUST_FACTORS["현장여건"])),
    ("공사규모", "보통", "", list(R.ADJUST_FACTORS["공사규모"])),
    ("지질조건", "양호(토사)", "지반 관련 공종에만 적용", list(R.ADJUST_FACTORS["지질조건"])),
    ("기상기후", "일반", "옥외·구조물 공정에만 적용", list(R.ADJUST_FACTORS["기상기후"])),
    ("작업일수 산정단위(일)", 1, "단위작업 작업일수 올림 단위: 1(1일 미만은 1일, 샘플보고서) 또는 0.1(검토엑셀)", ["1", "0.1"]),
    ("작업조 최대배수", 2, "공기조정 시 단위작업 작업조 상한 = 작업조 × 배수", None),
    ("공사여건 보정일수", 0, "제2장 2.(6) 공사여건 등에 따른 추가 공사기간(일)", None),
    ("공사여건 보정사유", "", "", None),
    ("■ 목표기간 및 공기조정", None, None, None),
    ("목표기간 기준", "없음", "없음 / 발주처요구 / 건설사업관리기간", ["없음", "발주처요구", "건설사업관리기간"]),
    ("발주처 요구 공사기간(개월)", None, "목표기간 기준이 '발주처요구'일 때", None),
    ("공기조정 결과 채택", "Y", "목표 초과 시 조정안을 최종 공사기간으로 채택", ["Y", "N"]),
    ("■ 적정성 검토 변수", None, None, None),
    ("허용범위(±)", 0.2, "가이드라인 ±20%", None),
    ("적정성 판정기준", "자동", "자동(부록5·동종실적 중 하나 이상 충족) / 부록5 / 동종공사 / 모두", ["자동", "부록5", "동종공사", "모두"]),
    ("연면적(100㎡)", None, "A — 기타건축물", None),
    ("지상층수", None, "G — 기타건축물·공동주택", None),
    ("지하층수", None, "B — 기타건축물", None),
    ("도로연장(m)", None, "L — 도로포장·도로", None),
    ("도로폭원(m)", None, "W — 도로(토공+교량)", None),
    ("교량연장(m)", None, "BL — 도로(토공+교량)", None),
    ("관경(mm)", None, "D — 상수도", None),
    ("양수·배수·가압장 개수", None, "S — 상수도", None),
    ("하수도연장(m)", None, "SL — 하수도", None),
    ("궤도연장(m)", None, "RL — 철도(궤도)", None),
    ("시설용량(㎥/일)", None, "V — 환경공단 산정식", None),
    ("환경공단 산정식", "해당없음", "환경시설 참고기준", list(R.KECO_FORMULAS) + ["해당없음"]),
    ("환경공단 경험치", "해당없음", "참고기준", list(R.KECO_EXPERIENCE) + ["해당없음"]),
    ("유사사례 준비·정리 가산", "Y", "가이드라인: 실적평균 + 준비·정리기간", ["Y", "N"]),
    ("■ 보고서 서술 (비우면 자동문구)", None, None, None),
    ("사업목적", "", "제1장 1.1", None),
    ("사업내용", "", "제1장 1.2 (줄바꿈으로 항목 구분)", None),
    ("추진경위", "", "제1장 (줄바꿈 구분)", None),
    ("공사특성", "", "제2장 2.2.1 (줄바꿈 구분)", None),
    ("시공조건 및 지연요인", "", "제2장 기타 공사기간 지연사항", None),
    ("검토자", "", "보고서 표지", None),
]
FORMULA_VAR_KEYS = {"연면적(100㎡)": "A", "지상층수": "G", "지하층수": "B", "도로연장(m)": "L", "도로폭원(m)": "W",
                    "교량연장(m)": "BL", "관경(mm)": "D", "양수·배수·가압장 개수": "S", "하수도연장(m)": "SL",
                    "궤도연장(m)": "RL", "시설용량(㎥/일)": "V"}

ITEM_COLS = ["No", "공정ID", "공종명(내역)", "규격", "단위", "수량", "DB지정(번호/공종명)", "수동1일작업량(품셈·실적)",
             "작업조", "최대작업조", "개별보정계수", "간접작업일수", "보정일수", "중복가능(Y)", "비고"]
ACT_COLS = ["공정ID", "분야", "대공정", "공정명", "기상프로필", "직접작업일수", "선행공정", "최대중첩률(0~1)",
            "가중치(공사비,백만원)", "마일스톤", "비고", "착수가능일", "소속공사"]
CASE_COLS = ["사업명", "시설용량", "공사금액(억원)", "공사기간(개월)", "준공년도", "비고"]

HFILL = PatternFill("solid", fgColor="1F4E78")
SFILL = PatternFill("solid", fgColor="DDEBF7")
YFILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass
class Project:
    ov: dict
    profiles: dict
    dust: list
    items: list
    acts: list
    cases: list
    path: str = ""
    warnings: list = field(default_factory=list)
    custom: list = field(default_factory=lambda: [0.0] * 12)
    custom_label: str = "사용자정의"
    holiday_off: set = field(default_factory=set)
    winter_exempt: set = field(default_factory=set)

    def g(self, k, default=None):
        v = self.ov.get(k)
        return default if v in (None, "") else v

    @property
    def name(self): return self.g("공사명", "(공사명 미입력)")
    @property
    def total_cost(self): return float(self.g("총공사비(백만원)", 0) or 0)
    @property
    def facility_type(self):
        t = self.g("시설유형(부록5)")
        return None if t in (None, "해당없음") else t

    @property
    def start_date(self):
        d = self.g("착공예정일")
        if isinstance(d, dt.datetime):
            return d.date()
        if isinstance(d, dt.date):
            return d
        return dt.date.fromisoformat(str(d)[:10])

    @property
    def prep_days(self):
        v = self.g("준비기간(일)")
        if v not in (None, ""):
            return int(v)
        return R.PREP_PERIODS.get(self.g("준비기간 공사유형"), 60)

    @property
    def cleanup_days(self): return int(self.g("정리기간(일)", 30))
    @property
    def tolerance(self): return float(self.g("허용범위(±)", 0.2))
    @property
    def winter_mode(self): return self.g("동절기 공사중지", "추가검토")
    @property
    def winter_months(self):
        return tuple(int(x) for x in str(self.g("동절기 중지월", "12,1,2")).replace(" ", "").split(",") if x)

    @property
    def formula_vars(self):
        out = {}
        for k, key in FORMULA_VAR_KEYS.items():
            v = self.g(k)
            if v not in (None, ""):
                out[key] = float(v)
        return out

    @property
    def keco_type(self):
        t = self.g("환경공단 산정식")
        return None if t in (None, "해당없음") else t

    @property
    def keco_exp(self):
        t = self.g("환경공단 경험치")
        return None if t in (None, "해당없음") else t

    @property
    def cases_add_prep(self): return str(self.g("유사사례 준비·정리 가산", "Y")).upper() == "Y"

    def factors(self):
        return Factors(self.g("현장여건", "보통"), self.g("공사규모", "보통"), self.g("지질조건", "양호(토사)"),
                       self.g("기상기후", "일반"))

    def target_days(self):
        mode = self.g("목표기간 기준", "없음")
        if mode == "발주처요구" and self.g("발주처 요구 공사기간(개월)"):
            m = float(self.g("발주처 요구 공사기간(개월)"))
            return round(m * R.DAYS_PER_MONTH), f"발주처 요구 공사기간 {m:g}개월"
        if mode == "건설사업관리기간":
            m = R.cm_period_months(self.total_cost / 100)
            return round(m * R.DAYS_PER_MONTH), f"건설사업관리기간 {m:.1f}개월(공사비 {self.total_cost / 100:,.1f}억원)"
        return None, None


# ── 작성 ──────────────────────────────────────────────────────────────────────
def _hdr(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c)
        cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), HFILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def write_template(path, ov=None, profiles=None, dust=None, items=None, acts=None, cases=None,
                   custom=None, custom_label="사용자정의", holiday_off=None, winter_exempt=None):
    """입력서식 작성. 값을 주면 예제 입력파일로 작성."""
    ov = ov or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "0_사용안내"
    guide = [
        "공사기간 적정성 검토 입력서식 — " + R.GUIDELINE,
        "",
        "작성 순서 (검토엑셀 ①→⑦ 순서와 동일)",
        "1_공사개요 : 사업정보, 착공일, 관측지점, 준비·정리기간, 보정계수, 목표기간, 부록5 변수, 보고서 서술",
        "2_비작업조건 : 기상프로필(공정 유형)별로 적용할 기상조건과 적용률(1=전부, 0.5=1/2) 선택, 미세먼지 월별 발령일수",
        "3_수량 : 설계내역 수량. 공정ID로 4_공정과 연결. 1일작업량은 수동입력 > DB지정 > 자동매칭 순으로 결정",
        "4_공정 : 토목·건축·기계·전기·시운전 등 공정과 선후관계(CPM). 수량이 없는 공정은 '직접작업일수' 입력",
        "5_유사사례 : 최근 5년 준공 동종·유사공사 (적정성 비교용)",
        "",
        "선행공정 표기: C10 (FS) / C10:FS+5 (완료 5일 후) / C10:SS50% (선행 50% 진행 후) / C10:SS+10 / C10:FF+0 — 쉼표로 여러 개",
        "최대중첩률: 공사기간이 목표를 초과할 때 해당 공정을 선행공정과 겹쳐 수행 가능한 최대 비율(0=불가, 0.5=50%)",
        "최대작업조: 비우면 '작업조 × 작업조 최대배수(1_공사개요)'까지 공기조정 시 증원 가능",
        "기상프로필: " + " / ".join(f"{k}({v})" for k, v in R.PROFILE_DESC.items()),
        "",
        "노란 셀이 입력 셀입니다. 실행: python3 -m cpa run <입력파일.xlsx> -o <출력폴더>",
    ]
    for i, t in enumerate(guide, 1):
        ws.cell(i, 1, t).font = Font(bold=(i in (1, 3)), size=12 if i == 1 else 10)
    ws.column_dimensions["A"].width = 140

    # 1_공사개요
    ws = wb.create_sheet("1_공사개요")
    _hdr(ws, 1, ["항목", "입력값", "설명"], [30, 48, 60])
    r = 2
    for key, default, desc, opts in OVERVIEW:
        if default is None and desc is None:
            c = ws.cell(r, 1, key)
            c.font, c.fill = Font(bold=True), SFILL
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        else:
            ws.cell(r, 1, key).border = BOX
            val = ov.get(key, default)
            c = ws.cell(r, 2, val)
            c.fill, c.border = YFILL, BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(val, (dt.date, dt.datetime)):
                c.number_format = "yyyy-mm-dd"
            ws.cell(r, 3, desc).border = BOX
            if opts:
                dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True) \
                    if len(",".join(opts)) < 250 else None
                if dv is None:  # 목록이 길면 숨김 목록 시트 사용
                    lst = wb["_목록"] if "_목록" in wb.sheetnames else wb.create_sheet("_목록")
                    col = lst.max_column + 1 if lst.max_row > 1 or lst["A1"].value else 1
                    for i, o in enumerate(opts, 1):
                        lst.cell(i, col, o)
                    L = get_column_letter(col)
                    dv = DataValidation(type="list", formula1=f"=_목록!${L}$1:${L}${len(opts)}", allow_blank=True)
                    lst.sheet_state = "hidden"
                ws.add_data_validation(dv)
                dv.add(c)
            if isinstance(val, str) and "\n" in val:
                ws.row_dimensions[r].height = min(15 * (val.count("\n") + 1), 200)
        r += 1

    # 2_비작업조건
    ws = wb.create_sheet("2_비작업조건")
    profiles = profiles or R.DEFAULT_PROFILES
    pnames = list(profiles)
    ws.cell(1, 1, "기상프로필별 기상조건 적용률 (1=전부 적용, 0.5=1/2 적용, 빈칸=미적용). 동절기는 결합조건(⑦) 1개, 강우는 공종 특성에 맞는 임계값 1개만 선택 권장(중복계상 방지)").font = Font(bold=True)
    _hdr(ws, 2, ["번호", "구분", "기상조건"] + pnames, [6, 10, 46] + [14] * len(pnames))
    for i, c in enumerate(range(1, 16)):
        rr = 3 + i
        ws.cell(rr, 1, c)
        ws.cell(rr, 2, R.COND_GROUP[c])
        ws.cell(rr, 3, R.CONDITIONS[c] if c != 15 else f"사용자정의:{custom_label} (아래 월별 입력)")
        for j, pn in enumerate(pnames):
            cell = ws.cell(rr, 4 + j, profiles[pn].get(c))
            cell.fill, cell.border = YFILL, BOX
    hoff, wex = holiday_off if holiday_off is not None else set(R.NO_HOLIDAY_PROFILES), winter_exempt or set()
    for k, (lab, st) in enumerate((("법정공휴일 적용", hoff), ("동절기 중지 적용", wex))):
        rr = 18 + k
        ws.cell(rr, 3, lab).font = Font(bold=True)
        for j, pn in enumerate(pnames):
            cell = ws.cell(rr, 4 + j, "N" if pn in st else "Y")
            cell.fill, cell.border = YFILL, BOX
    rr = 3 + 15 + 2
    ws.cell(rr, 1, "미세먼지 비상저감조치 월별 발령일수 (일) — 에어코리아 권역별 발령내역 등").font = Font(bold=True)
    _hdr(ws, rr + 1, ["", "", "월"] + [f"{m}월" for m in range(1, 13)])
    ws.cell(rr + 2, 3, "발령일수")
    for m in range(12):
        cell = ws.cell(rr + 2, 4 + m, (dust or R.DUST_EXAMPLES["미적용"])[m])
        cell.fill, cell.border = YFILL, BOX
    ws.cell(rr + 3, 3, "사용자정의")
    ws.cell(rr + 3, 2, custom_label)
    for m in range(12):
        cell = ws.cell(rr + 3, 4 + m, (custom or [0] * 12)[m])
        cell.fill, cell.border = YFILL, BOX
    ws.cell(rr + 5, 1, "※ 15번 사용자정의 행: 부록3에 없는 기상조건(예: 일최고 5℃ 이하) 또는 별도 산정한 월별 기상 비작업일수를 '사용자정의' 행에 입력(B열에 명칭).")
    ws.cell(rr + 6, 1, "※ 프로필을 추가하려면 열을 추가하고 머리글에 이름을 입력. '법정공휴일 적용=N'이면 공휴일·주40시간 하한 미적용(시운전 등 연속작업), '동절기 중지 적용=N'이면 동절기 공사중지 시에도 진행.")

    # 3_수량
    ws = wb.create_sheet("3_수량")
    _hdr(ws, 1, ITEM_COLS, [5, 8, 34, 26, 7, 11, 22, 14, 7, 9, 10, 9, 8, 9, 20])
    for i, it in enumerate(items or [], 2):
        for j, k in enumerate(ITEM_COLS, 1):
            ws.cell(i, j, it.get(k)).border = BOX
    for i in range(2, max(len(items or []) + 2, 60)):
        for j in range(2, 15):
            ws.cell(i, j).fill = YFILL
    ws.freeze_panes = "C2"

    # 4_공정
    ws = wb.create_sheet("4_공정")
    _hdr(ws, 1, ACT_COLS, [8, 8, 14, 30, 12, 10, 30, 12, 14, 18, 24, 12, 14])
    dvp = DataValidation(type="list", formula1='"' + ",".join(pnames) + '"', allow_blank=True)
    dvf = DataValidation(type="list", formula1='"' + ",".join(R.FIELDS) + '"', allow_blank=True)
    ws.add_data_validation(dvp)
    ws.add_data_validation(dvf)
    for i, a in enumerate(acts or [], 2):
        for j, k in enumerate(ACT_COLS, 1):
            c = ws.cell(i, j, a.get(k))
            c.border = BOX
            if isinstance(a.get(k), (dt.date, dt.datetime)):
                c.number_format = "yyyy-mm-dd"
    for i in range(2, max(len(acts or []) + 2, 40)):
        for j in range(1, 14):
            ws.cell(i, j).fill = YFILL
        dvp.add(ws.cell(i, 5))
        dvf.add(ws.cell(i, 2))
    ws.freeze_panes = "B2"

    # 5_유사사례
    ws = wb.create_sheet("5_유사사례")
    _hdr(ws, 1, CASE_COLS, [40, 14, 14, 14, 10, 30])
    for i, c in enumerate(cases or [], 2):
        for j, k in enumerate(CASE_COLS, 1):
            ws.cell(i, j, c.get(k)).border = BOX
    for i in range(2, 20):
        for j in range(1, 7):
            ws.cell(i, j).fill = YFILL
    wb.save(path)


# ── 읽기 ──────────────────────────────────────────────────────────────────────
def _rows(ws, cols):
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {k: (row[idx[k]] if k in idx and idx[k] < len(row) else None) for k in cols}
        if any(v not in (None, "") for v in d.values()):
            out.append(d)
    return out


def _num(v, default=None):
    if v in (None, ""):
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def read_project(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    warns = []
    ov = {}
    for row in wb["1_공사개요"].iter_rows(min_row=2, values_only=True):
        if row[0] and not str(row[0]).startswith("■"):
            ov[str(row[0]).strip()] = row[1]

    ws = wb["2_비작업조건"]
    hdr = [c.value for c in ws[2]]
    pnames = [h for h in hdr[3:] if h]
    profiles = {p: {} for p in pnames}
    dust = [0] * 12
    custom, custom_label, hoff, wex = [0] * 12, "사용자정의", set(R.NO_HOLIDAY_PROFILES) & set(pnames), set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[2] == "법정공휴일 적용":
            hoff = {p for j, p in enumerate(pnames) if str(row[3 + j] or "Y").upper().startswith("N")}
        if row[2] == "동절기 중지 적용":
            wex = {p for j, p in enumerate(pnames) if str(row[3 + j] or "Y").upper().startswith("N")}
        if row[2] == "사용자정의":
            custom = [_num(x, 0) for x in row[3:15]]
            custom_label = row[1] or "사용자정의"
        if isinstance(row[0], (int, float)) and 1 <= row[0] <= 15:
            for j, p in enumerate(pnames):
                v = _num(row[3 + j])
                if v:
                    profiles[p][int(row[0])] = v
        if row[2] == "발령일수":
            dust = [_num(x, 0) for x in row[3:15]]

    fac_mult = _num(ov.get("작업조 최대배수"), 2) or 1
    step = _num(ov.get("작업일수 산정단위(일)"), 1) or 1
    items = []
    for d in _rows(wb["3_수량"], ITEM_COLS):
        if not d["공정ID"]:
            continue
        crews = _num(d["작업조"], 1) or 1
        crews = int(crews) if crews == int(crews) else crews
        mx = _num(d["최대작업조"])
        items.append(Item(no=int(_num(d["No"], len(items) + 1)), act=str(d["공정ID"]).strip(),
                          name=str(d["공종명(내역)"] or ""), spec=str(d["규격"] or ""), unit=str(d["단위"] or ""),
                          qty=_num(d["수량"], 0), db_hint=d["DB지정(번호/공종명)"],
                          manual_rate=_num(d["수동1일작업량(품셈·실적)"]), crews=crews,
                          max_crews=mx if mx else max(crews, round(crews * fac_mult)),
                          item_factor=_num(d["개별보정계수"], 1.0), indirect=_num(d["간접작업일수"], 0),
                          correction=_num(d["보정일수"], 0),
                          parallel=str(d["중복가능(Y)"] or "").upper().startswith("Y"), note=str(d["비고"] or ""), step=step))
    acts = []
    for d in _rows(wb["4_공정"], ACT_COLS):
        if not d["공정ID"]:
            continue
        prof = d["기상프로필"] or "옥외공사"
        if prof not in profiles:
            warns.append(f"공정 {d['공정ID']}: 기상프로필 '{prof}' 미정의 → '옥외공사' 적용")
            prof = "옥외공사" if "옥외공사" in profiles else pnames[0]
        acts.append(Activity(id=str(d["공정ID"]).strip(), field=d["분야"] or "토목", group=d["대공정"] or "",
                             name=d["공정명"] or "", profile=prof, direct_days=_num(d["직접작업일수"], 0),
                             preds=parse_preds(d["선행공정"]), max_overlap=_num(d["최대중첩률(0~1)"], 0) or 0,
                             weight=_num(d["가중치(공사비,백만원)"]), milestone=d["마일스톤"] or "",
                             note=d["비고"] or "", snet=_date(d["착수가능일"]), project=d["소속공사"] or ""))
    ids = {a.id for a in acts}
    for it in items:
        if it.act not in ids:
            warns.append(f"수량 No.{it.no} '{it.name}': 공정ID '{it.act}'가 4_공정에 없음 → 제외")
    items = [it for it in items if it.act in ids]
    for it in items:
        it.group = next(a.group for a in acts if a.id == it.act)
    cases = [c for c in _rows(wb["5_유사사례"], CASE_COLS) if c["사업명"]]
    for c in cases:
        c["공사기간(개월)"] = _num(c["공사기간(개월)"])
    return Project(ov, profiles, dust, items, acts, cases, path, warns, custom, custom_label, hoff, wex)
