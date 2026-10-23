"""DOCX 작성 도우미 — 샘플보고서 양식(머리글: 사업명/장 제목, 바닥글: 쪽번호, [표 n-m] 캡션)."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x4E, 0x78)


def _font(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def _field(par, instr):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    r._element.append(f1)
    r._element.append(it)
    r._element.append(f2)
    _font(r, 9)


def _border_bottom(par, color="1F4E78", sz=12):
    pPr = par._element.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), str(sz))
    bt.set(qn("w:color"), color)
    b.append(bt)
    pPr.append(b)


class Doc:
    def __init__(self, project_name):
        self.d = Document()
        self.pname = project_name
        self.ch = 0
        self.avail = 16.6
        self.tn = 0
        self.fn = 0
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.3
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)

    # ── 머리글/바닥글 ──
    def set_header(self, section, right_text):
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run(self.pname)
        _font(r, 8.5, color=RGBColor(0x40, 0x40, 0x40))
        hp.add_run("\t\t")
        r2 = hp.add_run(right_text)
        _font(r2, 9, bold=True, color=NAVY)
        r2.italic = True
        _border_bottom(hp)
        fp = section.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("- ")
        _font(r, 9)
        _field(fp, "PAGE")
        r = fp.add_run(" -")
        _font(r, 9)

    def first_header(self, right_text):
        self.set_header(self.d.sections[0], right_text)

    def new_section(self, header_text, landscape=False):
        from docx.enum.section import WD_SECTION
        s = self.d.add_section(WD_SECTION.NEW_PAGE)
        self.avail = 26.7 if landscape else 16.6
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Cm(29.7), Cm(21)
            s.left_margin = s.right_margin = Cm(1.5)
            s.top_margin, s.bottom_margin = Cm(1.6), Cm(1.4)
        else:
            s.orientation = WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = Cm(21), Cm(29.7)
            s.left_margin = s.right_margin = Cm(2.2)
            s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
        self.set_header(s, header_text)
        return s

    # ── 본문 요소 ──
    def cover(self, title, sub, date_text, org, extra=None):
        for _ in range(6):
            self.d.add_paragraph()
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(title), 20, True)
        for _ in range(2):
            self.d.add_paragraph()
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(sub), 28, True, NAVY)
        if extra:
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(extra), 11, color=RGBColor(0x55, 0x55, 0x55))
        for _ in range(12):
            self.d.add_paragraph()
        for t, sz in ((date_text, 16), (org, 16)):
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(t), sz, True)

    def toc(self, entries):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run("< 목  차 >"), 16, True)
        self.d.add_paragraph()
        for lvl, text in entries:
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8 * lvl)
            p.paragraph_format.space_after = Pt(6 if lvl == 0 else 2)
            _font(p.add_run(text), 12 if lvl == 0 else 10.5, lvl == 0)

    def chapter(self, no, title, subs=(), title_page=True):
        """장 표지(샘플 양식) + 새 구역."""
        self.ch, self.tn, self.fn = no, 0, 0
        self.new_section(f"제 {no} 장 {title}")
        for _ in range(8):
            self.d.add_paragraph()
        p = self.d.add_paragraph()
        _font(p.add_run(f"제 {no} 장  {title}"), 24, True, NAVY)
        _border_bottom(p, sz=18)
        self.d.add_paragraph()
        for s in subs:
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            _font(p.add_run(s), 13)
        if not title_page:
            return
        p = self.d.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(f"제{no}장 {title}"), 16, True)

    def h2(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(text), 13, True, NAVY)

    def h3(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(text), 11.5, True)

    def h4(self, text):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run(text), 10.5, True)

    def para(self, text, indent=1.2, size=10, bold=False, italic=False, color=None):
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        _font(r, size, bold, color)
        r.italic = italic
        return p

    def bullet(self, text, indent=1.2, mark="-"):
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent + 0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        _font(p.add_run(f"{mark} {text}"), 10)
        return p

    def note(self, text, indent=0.8):
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent = Cm(indent)
        _font(p.add_run(text), 8.5, color=RGBColor(0x44, 0x44, 0x44))
        return p

    def boxed(self, lines, title=None, fill="F2F2F2"):
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        _shade(c, fill)
        c.paragraphs[0].text = ""
        if title:
            _font(c.paragraphs[0].add_run(title), 10, True, NAVY)
            first = False
        else:
            first = True
        for ln in lines:
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(lines) == 1 else WD_ALIGN_PARAGRAPH.LEFT
            _font(p.add_run(ln), 10.5 if len(lines) == 1 else 9.5, len(lines) == 1)
        self.d.add_paragraph()

    def caption_table(self, title):
        self.tn += 1
        p = self.d.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        _font(p.add_run(f"[표 {self.ch}-{self.tn}]"), 9.5, True)
        _font(p.add_run(f"   {title}"), 10, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, header, rows, title=None, widths=None, size=8.5, align=None, hdr_fill="D9E2F3", bold_rows=(), fills=None, header_rows=None):
        if title:
            self.caption_table(title)
        hdrs = header_rows or [header]
        ncol = len(hdrs[0])
        t = self.d.add_table(rows=len(hdrs) + len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for hi, hr in enumerate(hdrs):
            for j, h in enumerate(hr):
                c = t.cell(hi, j)
                c.text = ""
                _shade(c, hdr_fill)
                pr = c.paragraphs[0]
                pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _font(pr.add_run(str(h)), size, True)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = t.cell(len(hdrs) + i, j)
                c.text = ""
                pr = c.paragraphs[0]
                pr.paragraph_format.space_after = Pt(0)
                pr.paragraph_format.line_spacing = 1.05
                txt = "" if v is None else (f"{v:,.1f}" if isinstance(v, float) and abs(v - round(v)) > 1e-9 else
                                             (f"{int(round(v)):,}" if isinstance(v, (int, float)) else str(v)))
                a = (align[j] if align else ("L" if isinstance(v, str) and len(v) > 14 else "C"))
                pr.alignment = {"L": WD_ALIGN_PARAGRAPH.LEFT, "C": WD_ALIGN_PARAGRAPH.CENTER, "R": WD_ALIGN_PARAGRAPH.RIGHT}[a]
                _font(pr.add_run(txt), size, i in bold_rows)
                if fills and fills[i]:
                    _shade(c, fills[i])
        self._widths(t, widths, ncol)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def _widths(self, t, widths, ncol):
        """열 너비 지정(Word·LibreOffice 공통: tblGrid + 셀 너비), 페이지 폭에 맞게 비례 조정."""
        if not widths:
            # 내용 길이에 비례한 기본 너비
            lens = [1] * ncol
            for row in t.rows:
                for j, c in enumerate(row.cells[:ncol]):
                    lens[j] = max(lens[j], min(len(c.text), 30))
            widths = [max(1.2, l) for l in lens]
        tot = sum(widths)
        widths = [w * self.avail / tot for w in widths] if tot > self.avail or not widths else widths
        t.autofit = False
        tbl = t._tbl
        grid = tbl.tblGrid
        for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
            if j < len(widths):
                gc.set(qn("w:w"), str(int(Cm(widths[j]).twips)))
        for row in t.rows:
            for j, c in enumerate(row.cells):
                if j < len(widths):
                    c.width = Cm(widths[j])

    def merge(self, t, r0, c0, r1, c1):
        a, b = t.cell(r0, c0), t.cell(r1, c1)
        txt = a.text
        m = a.merge(b)
        for p in m.paragraphs[1:]:
            p._element.getparent().remove(p._element)
        m.paragraphs[0].text = ""
        _font(m.paragraphs[0].add_run(txt), 8.5, True)
        m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return m

    def figure(self, path, title, width_cm=16, max_h_cm=None):
        if max_h_cm:
            from PIL import Image
            w, h = Image.open(path).size
            width_cm = min(width_cm, max_h_cm * w / h)
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Cm(width_cm))
        self.fn += 1
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(c.add_run(f"[그림 {self.ch}-{self.fn}] {title}"), 9.5, True)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.d.save(path)
