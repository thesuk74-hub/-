"""수량첨부 검토 — 공정 선택 + 공종별 수량파일 → 입력서식 자동 작성

  ① 공정 선택(처리시설공사: 처리시설공사·구내배관공사 / 관로공사: 관공사·배수설비 공사·펌프장 공사)
     공정별 다수의 수량파일(토공·구조물공·가시설공·포장공·부대공·관로공·기타 …)
  ② 2.1 공사개요 · 2.2 비작업조건(관측지점·기상조건 자동/선택)
     2.3 수량 ↔ 부록4 1일작업량 매칭 → 순작업일수 + 간접작업일수 + 보정일수, 공사개요 공사기간 초과 시 작업조·중첩 조정
     2.5 유사사례·유사공정(산정식) 선택
  → project.write_template() 입력서식 → pipeline.run() (①~⑦ 동일 엔진)

state(dict, meta["wiz"])
  works   : [{"name": 공정명, "files": [{"file", "orig", "category", "sheets_off": []}]}]
  ov      : 공사개요 값(1_공사개요 항목명)
  nonwork : {"auto": bool, "profiles": {프로필: {조건번호: 적용률}}, "dust": [12], "custom": [12], "custom_label"}
  items   : {행키: {"use", "db", "rate", "crews", "indirect", "correction", "act"}}
  acts    : {공정ID: {"direct", "overlap", "use"}}
  corr    : {공정명: [보정항목...]}, "corr_custom": {공정ID: 일수}
  indirect_auto : bool
  cases   : [유사사례]
"""
import datetime as dt
import json
import os
import re

from . import refdata as R
from .project import write_template
from .qtyparse import CATEGORY_AUTO, parse_file, row_category
from .workdays import DB_BY_NO, auto_match, norm_unit

# ── 공정 템플릿 ───────────────────────────────────────────────────────────────
# key: (공정명, 분야, 대공정, 기상프로필, 선행[구조물형], 선행[선형], 최대중첩률)
ACT_TPL = {
    "PV1": ("포장 절단·깨기", "토목", "포장공", "옥외공사", [], [], 0.3),
    "TS1": ("가시설(흙막이) 설치", "토목", "가시설공", "옥외공사", ["PV1"], ["PV1:SS10%"], 0.3),
    "EX": ("터파기", "토목", "토공", "옥외공사", ["TS1:SS50%", "PV1"], ["TS1:SS10%", "PV1:SS10%"], 0.3),
    "ST": ("구조물공", "토목", "구조물공", "구조물공사", ["EX"], ["PI:SS20%", "PI:FF+0"], 0.3),
    "PI": ("관로공(관 부설·접합)", "토목", "관로공", "옥외공사", ["ST:SS60%"], ["EX:SS10%", "EX:FF+0"], 0.3),
    "BF": ("되메우기", "토목", "토공", "옥외공사", ["ST"], ["PI:SS10%", "PI:FF+0"], 0.3),
    "TS2": ("가시설 해체", "토목", "가시설공", "옥외공사", ["BF:SS30%"], ["BF:SS10%", "BF:FF+0"], 0.3),
    "PV2": ("포장 복구", "토목", "포장공", "옥외공사", ["BF", "TS2"], ["BF:SS20%", "BF:FF+0", "TS2:FF+0"], 0.3),
    "AC": ("부대공", "토목", "부대공", "옥외공사", ["BF:SS50%"], ["PV2:SS50%", "PV2:FF+0"], 0.3),
    "AR": ("건축공사", "건축", "건축공사", "옥내공사", ["ST"], ["PI:SS50%"], 0.3),
    "ME": ("기계설비 설치", "기계", "기계공사", "옥내공사", ["ST", "AR:SS50%"], ["PI:SS50%"], 0.3),
    "EL": ("전기·계측제어 설치", "전기", "전기공사", "옥내공사", ["ME:SS50%", "ST"], ["ME:SS50%", "PI:SS50%"], 0.3),
    "ET": ("기타공", "토목", "기타", "옥외공사", ["ST:SS50%"], ["PI:SS50%"], 0.3),
    "TR": ("시운전", "시운전", "시운전", "연속작업", ["ME", "EL", "PV2", "PI", "AC", "ET"], ["ME", "EL"], 0.0),
}
ORDER = ["PV1", "TS1", "EX", "ST", "PI", "BF", "TS2", "PV2", "AC", "AR", "ME", "EL", "ET", "TR"]
CAT_ACT = {"토공": "EX", "가시설공": "TS1", "구조물공": "ST", "관로공": "PI", "포장공": "PV2", "부대공": "AC",
           "건축": "AR", "기계": "ME", "전기": "EL", "기타": "ET"}
SPLIT = {  # 분류 내 세분: (정규식, 공정키)
    "토공": [(r"되메우기|성토|다짐|모래\s*(기초|부설)|되메움|채움", "BF")],
    "가시설공": [(r"해체|철거|인발|뽑기|회수", "TS2")],
    "포장공": [(r"깨기|절단|절삭|철거|제거|커팅", "PV1")],
}
PARALLEL_RE = re.compile(r"운반|사토|잔토|상차|물\s*푸기|양수|보양|청소|시험|측량|안전|교통|품질|정리|적치|폐기물|경비|손료", re.I)


def act_key(cat, name):
    for pat, k in SPLIT.get(cat, []):
        if re.search(pat, name):
            return k
    return CAT_ACT.get(cat, "ET")


def default_state():
    return {"works": [], "ov": {}, "nonwork": {"auto": True}, "items": {}, "acts": {}, "corr": {}, "corr_custom": {},
            "indirect_auto": True, "cases": [], "step": 1}


# ── 관측지점 자동선택 ─────────────────────────────────────────────────────────
REGION_ALIAS = {"강원": "강원도", "경기": "경기도", "경남": "경상남도", "경북": "경상북도", "광주": "광주광역시",
                "대구": "대구광역시", "대전": "대전광역시", "부산": "부산광역시", "서울": "서울특별시", "세종": "세종특별시",
                "울산": "울산광역시", "인천": "인천광역시", "전남": "전라남도", "전북": "전라북도", "제주": "제주특별시",
                "충남": "충청남도", "충북": "충청북도", "경상남": "경상남도", "경상북": "경상북도", "전라남": "전라남도",
                "전라북": "전라북도", "충청남": "충청남도", "충청북": "충청북도"}
REGION_DEFAULT = {"강원도": "춘천", "경기도": "수원", "경상남도": "창원", "경상북도": "안동", "인천광역시": "인천",
                  "전라남도": "목포", "전라북도": "전주", "제주특별시": "제주", "충청남도": "홍성", "충청북도": "청주"}


def auto_station(loc):
    """사업위치(지역) 문자열 → (관측지점, 사유)."""
    loc = str(loc or "")
    st = R.WEATHER["stations"]
    for name in sorted(st, key=len, reverse=True):
        base = re.sub(r"(시|군)$", "", name)
        if len(base) >= 2 and base in loc:
            return name, f"사업위치에 '{base}' 포함 → 관측지점 '{name}'"
    for k in sorted(REGION_ALIAS, key=len, reverse=True):
        if k in loc:
            reg = REGION_ALIAS[k]
            s = REGION_DEFAULT.get(reg) or next(n for n, v in st.items() if v["region"] == reg)
            return s, f"시·군 관측지점 없음 → {reg} 대표 관측지점 '{s}'(최인접 지점 확인 필요)"
    return "서울", "사업위치에서 지역을 찾지 못해 '서울' 적용(관측지점 선택 필요)"


def auto_dust(station):
    reg = R.WEATHER["stations"].get(station, {}).get("region", "")
    if reg in ("서울특별시", "경기도", "인천광역시"):
        return "수도권(가이드라인 연평균 약 5일)"
    if reg == "충청남도":
        return "충남서부권역(서산 사례)"
    return "미적용"


# ── 수량 행 → 단위작업 ────────────────────────────────────────────────────────
def parsed(base_dir, f):
    """파일 해석 결과(캐시)."""
    p = os.path.join(base_dir, f["file"])
    cache = p + ".json"
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(p):
        with open(cache, encoding="utf-8") as fh:
            return json.load(fh)
    r = parse_file(p)
    with open(cache, "w", encoding="utf-8") as fh:
        json.dump(r, fh, ensure_ascii=False)
    return r


def _indirect(text):
    for pat, what, days in R.INDIRECT_RULES:
        if re.search(pat, text):
            return pat, what, days
    return None


def collect(state, base_dir):
    """모든 수량 행을 단위작업 후보로 정리. 반환: [row dict] (공정ID·매칭·작업일수 미리보기 포함)."""
    rows = []
    ov = state.get("ov", {})
    for w in state["works"]:
        info = R.WORK_INFO.get(w["name"], {"code": "X", "type": "구조물형"})
        for fi, f in enumerate(w["files"]):
            res = parsed(base_dir, f)
            for sh in res["sheets"]:
                if sh["sheet"] in (f.get("sheets_off") or []):
                    continue
                for r in sh["rows"]:
                    key = f"{w['name']}|{f['file']}|{sh['sheet']}|{r['row']}"
                    cat = row_category(f.get("category"), sh["sheet"], r)
                    k = act_key(cat, f"{r['name']} {r['spec']}")
                    o = state["items"].get(key, {})
                    act = o.get("act") or f"{info['code']}-{k}"
                    field = ACT_TPL.get(act.split("-", 1)[-1], ACT_TPL["ET"])[1]
                    db, note, conf = None, "", 0.0
                    if o.get("db"):
                        db = DB_BY_NO.get(int(o["db"]))
                        note = "DB지정(검토자)"
                    elif not o.get("rate"):
                        db, note, conf = auto_match(r["name"], r["spec"], r["unit"], field)
                    par_auto = bool(PARALLEL_RE.search(r["name"])) and not o.get("db") and not o.get("rate")
                    matched = bool(db) or bool(o.get("rate"))
                    use = o.get("use", matched and not par_auto)
                    rows.append({"key": key, "work": w["name"], "file": f.get("orig") or os.path.basename(f["file"]),
                                 "sheet": sh["sheet"], "row": r["row"], "category": cat, "act": act,
                                 "name": r["name"], "spec": r["spec"], "unit": r["unit"], "qty": r["qty"],
                                 "db": db["no"] if db else None, "db_name": f"{db['name']} ({db['condition']})" if db else "",
                                 "db_rate": db["rate"] if db else None, "db_unit": db["unit"] if db else "",
                                 "rate": o.get("rate"), "crews": o.get("crews") or 1, "use": bool(use),
                                 "match": note if (db or o.get("rate")) else note or "매칭불가",
                                 "conf": conf, "unit_ok": (not db) or norm_unit(db["unit"]) == norm_unit(r["unit"]),
                                 "excluded_reason": "" if use else ("동시작업(운반·시험 등) — 공기 미합산" if par_auto and o.get("use") is None
                                                                    else ("1일작업량 미매칭 — DB지정 또는 수동입력 필요" if not matched else "검토자 제외")),
                                 "indirect": o.get("indirect"), "correction": o.get("correction") or 0})
    # 간접작업일수 자동(공정 내 동일 단위공정 1회)
    seen = set()
    for x in rows:
        x["indirect_auto"], x["indirect_note"] = 0.0, ""
        if not x["use"] or not state.get("indirect_auto", True):
            continue
        hit = _indirect(f"{x['db_name']} {x['name']}")
        if hit and (x["act"], hit[0]) not in seen:
            seen.add((x["act"], hit[0]))
            x["indirect_auto"], x["indirect_note"] = hit[2], hit[1]
    for x in rows:
        x["indirect_applied"] = x["indirect"] if x["indirect"] not in (None, "") else x["indirect_auto"]
        rate = x["rate"] or x["db_rate"]
        x["days"] = None
        if x["use"] and rate:
            import math
            x["days"] = max(1, math.ceil(x["qty"] / (float(rate) * float(x["crews"] or 1)) - 1e-9))
    return rows


def activities(state, rows):
    """공정 구성(템플릿) — 수량이 있는 공정 + 시운전. 반환: [act dict]"""
    acts = []
    present = {}
    for x in rows:
        present.setdefault(x["act"], []).append(x)
    by_work = {}
    for w in state["works"]:
        info = R.WORK_INFO.get(w["name"], {"code": "X", "type": "구조물형", "commission": 0})
        code, typ = info["code"], info["type"]
        keys = [k for k in ORDER if f"{code}-{k}" in present]
        comm = state["acts"].get(f"{code}-TR", {}).get("direct", info.get("commission", 0))
        if comm and typ == "구조물형":
            keys.append("TR")
        keys = [k for k in keys if state["acts"].get(f"{code}-{k}", {}).get("use", True)]
        ks = set(keys)
        tpl_i = 4 if typ == "구조물형" else 5

        def resolve(tok, depth=0):
            k, _, rel = tok.partition(":")
            if k in ks:
                return [tok]
            if depth > 6 or k not in ACT_TPL:
                return []
            out = []
            for t in ACT_TPL[k][tpl_i]:
                out += resolve(t, depth + 1)
            return out

        wacts = []
        for k in keys:
            nm, field, grp, prof, _, _, ov = ACT_TPL[k]
            preds = []
            seen = set()
            for t in ACT_TPL[k][tpl_i]:
                for r in resolve(t):
                    pid, _, rel = r.partition(":")
                    kind = "FF" if rel.upper().startswith("FF") else "S"
                    if pid != k and (pid, kind) not in seen:
                        seen.add((pid, kind))
                        preds.append(r)
            aid = f"{code}-{k}"
            o = state["acts"].get(aid, {})
            direct = o.get("direct", comm if k == "TR" else 0) or 0
            wacts.append({"공정ID": aid, "분야": field, "대공정": grp, "공정명": f"[{w['name']}] {nm}", "기상프로필": prof,
                          "직접작업일수": direct, "선행공정": [f"{code}-{p}" for p in preds],
                          "최대중첩률(0~1)": o.get("overlap", ov), "소속공사": w["name"], "key": k,
                          "비고": "시운전 기간(검토자 입력)" if k == "TR" else ""})
        by_work[w["name"]] = (code, wacts)
        acts += wacts

    def first(wname):
        return [a for a in by_work.get(wname, ("", []))[1] if not a["선행공정"]]

    def last(wname):
        ws = by_work.get(wname, ("", []))[1]
        preds = {p.split(":")[0] for a in ws for p in a["선행공정"]}
        return [a for a in ws if a["공정ID"] not in preds]

    def anchor(wname, keys):
        code, ws = by_work.get(wname, ("", []))
        ids = {a["key"]: a["공정ID"] for a in ws}
        for k in keys:
            if k in ids:
                return ids[k]
        return None

    # 공정 간 연계: 구내배관 ← 처리시설 구조물 SS50%, 처리시설 시운전 ← 구내배관 완료, 배수설비 ← 관공사 관부설 SS30%
    t_st = anchor("처리시설공사", ["ST", "EX", "AR"])
    if t_st:
        for a in first("구내배관공사"):
            a["선행공정"].append(f"{t_st}:SS50%")
        t_tr = anchor("처리시설공사", ["TR"])
        if t_tr:
            ta = next(a for a in acts if a["공정ID"] == t_tr)
            ta["선행공정"] += [a["공정ID"] for a in last("구내배관공사")]
    p_pi = anchor("관공사", ["PI", "EX"])
    if p_pi:
        for a in first("배수설비 공사"):
            a["선행공정"].append(f"{p_pi}:SS30%")
    for a in acts:
        a["선행공정"] = ", ".join(a["선행공정"])
    return acts


# ── 입력서식 작성 ─────────────────────────────────────────────────────────────
def corrections(state, acts):
    """보정일수 항목 → [(공정ID, 항목, 일수, 사유)]"""
    out = []
    ids = {a["공정ID"] for a in acts}
    for w in state["works"]:
        code = R.WORK_INFO.get(w["name"], {"code": "X"})["code"]
        for item in state["corr"].get(w["name"], []):
            days, target, desc = R.CORRECTION_ITEMS[item]
            keys = ["TS1", "EX", "ST", "PI"] if target.startswith("가시설") else ["ST", "EX", "PI"]
            aid = next((f"{code}-{k}" for k in keys if f"{code}-{k}" in ids), None)
            if aid:
                out.append((aid, item, days, desc))
    for aid, d in (state.get("corr_custom") or {}).items():
        if aid in ids and d:
            out.append((aid, "공사여건 보정(검토자)", float(d), "제2장 2.(6) 공사여건 등에 따른 보정"))
    return out


def build_input(state, base_dir, path):
    rows = collect(state, base_dir)
    acts = activities(state, rows)
    ids = {a["공정ID"] for a in acts}
    items, excluded = [], []
    for x in rows:
        if x["act"] not in ids:
            continue
        if not x["use"] and not str(x["excluded_reason"]).startswith("동시작업"):
            excluded.append(x)          # 미매칭·검토자 제외 → 작업일수 미반영(제외수량 시트에 기록)
            continue
        items.append({"No": len(items) + 1, "공정ID": x["act"], "공종명(내역)": x["name"], "규격": x["spec"], "단위": x["unit"],
                      "수량": x["qty"], "DB지정(번호/공종명)": x["db"] if not x["rate"] else None,
                      "수동1일작업량(품셈·실적)": x["rate"], "작업조": x["crews"], "간접작업일수": x["indirect_applied"] if x["use"] else None,
                      "보정일수": x["correction"] if x["use"] else None, "중복가능(Y)": None if x["use"] else "Y",
                      "비고": " / ".join(t for t in (f"{x['file']}·{x['sheet']}·{x['row']}행", x["indirect_note"] and f"간접:{x['indirect_note']}",
                                                     x["excluded_reason"]) if t)})
    for aid, item, days, desc in corrections(state, acts):
        items.append({"No": len(items) + 1, "공정ID": aid, "공종명(내역)": f"[보정일수] {item}", "규격": desc, "단위": "식",
                      "수량": 0, "수동1일작업량(품셈·실적)": 1, "작업조": 1, "보정일수": days,
                      "비고": R.CORRECTION_SOURCE})
    ov = overview(state, rows)
    nw = state.get("nonwork") or {}
    profiles = {k: {int(c): float(v) for c, v in d.items()} for k, d in (nw.get("profiles") or R.DEFAULT_PROFILES).items()}
    dust = nw.get("dust") or R.DUST_EXAMPLES[auto_dust(ov["기상관측지점"])]
    cases = [{k: c.get(k) for k in ("사업명", "시설용량", "공사금액(억원)", "공사기간(개월)", "준공년도", "비고")}
             for c in state.get("cases") or [] if c.get("사업명")]
    acts_out = [{k: v for k, v in a.items() if k != "key"} for a in acts]
    write_template(path, ov, profiles, dust, items, acts_out, cases, custom=nw.get("custom"),
                   custom_label=nw.get("custom_label") or "사용자정의")
    _write_excluded(path, excluded)
    return {"rows": rows, "acts": acts, "ov": ov, "excluded": excluded}


def _write_excluded(path, excluded):
    """검토입력.xlsx 에 '6_제외수량' 시트(작업일수 미반영 수량과 사유) 추가."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("6_제외수량")
    hdr = ["공정", "공정ID", "출처(파일)", "시트", "행", "공종명", "규격", "단위", "수량", "사유"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for i, x in enumerate(excluded, 2):
        for j, v in enumerate([x["work"], x["act"], x["file"], x["sheet"], x["row"], x["name"], x["spec"], x["unit"], x["qty"],
                               x["excluded_reason"] or x["match"]], 1):
            ws.cell(i, j, v)
    for col, w in zip("ABCDEFGHIJ", (14, 9, 18, 12, 6, 30, 26, 7, 10, 44)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def pipe_length(rows):
    """하수도연장(SL) 추정: 관로공 m 단위 관 부설 수량 합."""
    codes = tuple(f"{v['code']}-" for v in R.WORK_INFO.values() if v["group"] == "관로공사")
    return round(sum(x["qty"] for x in rows if x["act"].endswith("-PI") and x["act"].startswith(codes) and norm_unit(x["unit"]) == "m"
                     and re.search(r"관|부설|접합", x["name"]) and not re.search(r"절단|접합부|보호", x["name"])), 1)


def suggest_formulas(state, rows=None):
    """유사공정(산정식) 추천 — 샘플보고서(대죽): 부록5 하수도 / 환경공단 산정식(폐수처리시설) / 경험치."""
    works = {w["name"] for w in state["works"]}
    groups = {R.WORK_INFO[w]["group"] for w in works if w in R.WORK_INFO}
    s = {"시설유형(부록5)": "해당없음", "환경공단 산정식": "해당없음", "환경공단 경험치": "해당없음"}
    if "관로공사" in groups:
        s["시설유형(부록5)"] = "하수도"
        s["환경공단 경험치"] = "하수관로"
        if rows is not None:
            s["하수도연장(m)"] = pipe_length(rows) or None
    if "처리시설공사" in groups:
        s["환경공단 산정식"] = "폐수처리시설"
        s["환경공단 경험치"] = "처리장"
    return s


def overview(state, rows):
    o = dict(state.get("ov") or {})
    if not o.get("기상관측지점"):
        o["기상관측지점"] = auto_station(o.get("사업위치"))[0]
    for k, v in suggest_formulas(state, rows).items():
        o.setdefault(k, v)
    if o.get("준공예정일") and not o.get("목표기간 기준"):
        o["목표기간 기준"] = "준공예정일"
    if not o.get("공사명"):
        o["공사명"] = o.get("사업명") or "공사명 미입력"
    if not o.get("공사특성"):
        o["공사특성"] = "공사 구성: " + ", ".join(w["name"] for w in state["works"]) + \
                     "\n수량자료: " + ", ".join(f"{w['name']} {len(w['files'])}개 파일" for w in state["works"])
    for k in ("착공예정일", "준공예정일"):
        if isinstance(o.get(k), str) and o[k]:
            o[k] = dt.date.fromisoformat(o[k][:10])
    for k, v in list(o.items()):
        if isinstance(v, str) and re.fullmatch(r"-?\d+(\.\d+)?", v.strip()):
            o[k] = float(v) if "." in v else int(v)
    return o
