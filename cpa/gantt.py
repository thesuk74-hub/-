"""주공정(Critical Path) 공정표 — 샘플보고서 3.2 예정공정표 양식.

 · 주공정 : 녹색 막대 + 적색 화살표, 주공정 연결 적색 점선
 · 비주공정 : 흑색 화살표,  총여유(TF) : 회색 막대
 · 마일스톤, 준비·정리기간, 월별 계획 공정률(%)
"""
import calendar
import datetime as dt

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402
from matplotlib.patches import FancyArrowPatch, Rectangle  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from . import refdata as R  # noqa: E402
from .summary import month_iter, progress_curve  # noqa: E402

for fam in ("Malgun Gothic", "NanumGothic", "NanumBarunGothic", "AppleGothic", "WenQuanYi Zen Hei"):
    if any(fam == f.name for f in font_manager.fontManager.ttflist):
        plt.rcParams["font.family"] = fam
        break
plt.rcParams["axes.unicode_minus"] = False

GREEN, RED, GREY, BLUE, NAVY = "#8BC34A", "#E53935", "#BDBDBD", "#1565C0", "#1F4E78"
FIELD_ORDER = ["토목", "건축", "기계", "전기", "통신", "소방", "조경", "기타", "시운전"]


def chart_rows(res):
    """공정표 행 구성: (종류, 라벨, 공정 or None)."""
    rows = [("milestone", "Mile Stone", None), ("prep", "준비기간", None)]
    acts = [a for a in res.acts]
    fields = [f for f in FIELD_ORDER if any(a.field == f for a in acts)] + \
             sorted({a.field for a in acts} - set(FIELD_ORDER))
    for f in fields:
        fa = sorted([a for a in acts if a.field == f], key=lambda a: (a.es, a.ef))
        if f == "시운전" and len(fa) == 1:
            rows.append(("act", "■ " + fa[0].name, fa[0]))
            continue
        rows.append(("field", f"■ {f}", fa))
        for a in fa:
            rows.append(("act", f"  ▶ {a.name}", a))
    rows.append(("clean", "정리기간", None))
    return rows


def _x(cal, day):
    return day


def render(res, p, path_png, path_pdf=None, title=None):
    cal = res.cal
    total = res.total_days
    start, end = cal.day_to_date(0), cal.day_to_date(total - 1)
    months = list(month_iter(start, end))
    rows = chart_rows(res)
    prog = progress_curve(res)
    n = len(rows)
    fig_h = max(8.3, 2.6 + 0.42 * (n + 2))
    fig = plt.figure(figsize=(16.54, fig_h))
    top = 1 - 1.75 / fig_h
    ax = fig.add_axes([0.17, 0.03, 0.76, top - 0.03])
    # x축 = 달력일차, 월 경계
    def mday(y, m):
        return cal.date_to_day(dt.date(y, m, 1))
    x0, x1 = mday(*months[0]), mday(*months[-1]) + calendar.monthrange(*months[-1])[1]
    ax.set_xlim(x0, x1)
    ax.set_ylim(n + 2, -1.6)
    ax.axis("off")
    # 월 그리드
    for y, m in months:
        xs = mday(y, m)
        ax.axvline(xs, color="#CCCCCC", lw=0.6, zorder=0)
        ax.text(xs + calendar.monthrange(y, m)[1] / 2, -0.6, f"{m}월", ha="center", va="center", fontsize=8)
    years = {}
    for y, m in months:
        years.setdefault(y, []).append(m)
    for y, ms in years.items():
        a, b = mday(y, ms[0]), mday(y, ms[-1]) + calendar.monthrange(y, ms[-1])[1]
        ax.text((a + b) / 2, -1.25, f"{y}년", ha="center", va="center", fontsize=9, fontweight="bold")
        ax.plot([a, a], [-1.55, n + 2], color="#888888", lw=0.9)
    ax.plot([x0, x1], [-0.25, -0.25], color="#555555", lw=1)
    ax.plot([x0, x1], [-1.55, -1.55], color="#555555", lw=1)
    # 좌측 라벨·우측 기간
    for i, (kind, label, obj) in enumerate(rows):
        y = i + 0.5
        ax.axhline(i + 1, color="#E0E0E0", lw=0.5, zorder=0)
        ax.text(x0 - (x1 - x0) * 0.005, y, label, ha="right", va="center", fontsize=8.5,
                fontweight="bold" if kind in ("field", "milestone") else "normal", clip_on=False)
        dur = None
        if kind == "prep":
            dur = res.prep
        elif kind == "clean":
            dur = res.cleanup
        elif kind == "act" and obj.work > 0:
            dur = obj.cal_days
        elif kind == "field":
            es, ef = min(a.es for a in obj), max(a.ef for a in obj)
            dur = ef - es + 1
        if dur:
            txt = f"{dur / R.DAYS_PER_MONTH:.1f}개월"
            if kind == "act" and not obj.critical:
                txt = f"({txt})"
            ax.text(x1 + (x1 - x0) * 0.006, y, txt, ha="left", va="center", fontsize=8, clip_on=False)
    ax.text(x1 + (x1 - x0) * 0.006, -0.9, "Calendar\nDay", ha="left", va="center", fontsize=8, fontweight="bold", clip_on=False)
    ax.text(x0 - (x1 - x0) * 0.08, -0.9, "구 분", ha="center", va="center", fontsize=10, fontweight="bold", clip_on=False)

    def bar(xa, xb, y, crit, tf_to=None):
        if crit:
            ax.add_patch(Rectangle((xa, y - 0.17), xb - xa + 1, 0.34, color=GREEN, zorder=2))
            ax.add_patch(FancyArrowPatch((xa, y), (xb + 1, y), arrowstyle="-|>", mutation_scale=9, color=RED, lw=1.6, zorder=3))
            ax.plot([xa], [y], "o", color=RED, ms=3.5, zorder=4)
        else:
            ax.add_patch(FancyArrowPatch((xa, y), (xb + 1, y), arrowstyle="-|>", mutation_scale=8, color="black", lw=1.1, zorder=3))
            ax.plot([xa], [y], "o", color="black", ms=3, zorder=4)
        if tf_to and tf_to > xb:
            ax.add_patch(Rectangle((xb + 1, y - 0.13), tf_to - xb, 0.26, color=GREY, zorder=1))

    ypos = {}
    for i, (kind, label, obj) in enumerate(rows):
        y = i + 0.5
        if kind == "prep":
            bar(0, res.prep_end, y, True)
            ax.text((res.prep_end + 1) / 2, y - 0.32, "준비기간", ha="center", fontsize=7.5, color=RED, fontweight="bold")
        elif kind == "clean":
            bar(res.finish_day + 1, total - 1, y, True)
            ax.text(res.finish_day + res.cleanup / 2, y - 0.32, "정리기간", ha="center", fontsize=7.5, color=RED, fontweight="bold")
        elif kind == "act" and obj.work > 0:
            bar(obj.es, obj.ef, y, obj.critical, obj.lf if obj.tf > 0 else None)
            ax.text((obj.es + obj.ef) / 2, y - 0.3, obj.name, ha="center", va="bottom", fontsize=7, zorder=5,
                    bbox=dict(boxstyle="square,pad=0.1", fc="white", ec="none", alpha=0.7))
            ypos[obj.id] = y
        elif kind == "field":
            es, ef = min(a.es for a in obj), max(a.ef for a in obj)
            ax.add_patch(Rectangle((es, y - 0.08), ef - es + 1, 0.16, color=NAVY, alpha=0.35, zorder=1))
        elif kind == "milestone":
            ms = [("준비기간", res.prep_end)]
            for a in res.acts:
                if a.milestone:
                    ms.append((a.milestone, a.ef))
            ms.append(("정리기간", total - 1))
            for k, (nm, d) in enumerate(ms):
                col = RED if nm in ("준비기간", "정리기간") else "#2F6EB5"
                ax.annotate(nm, (d, y + 0.25), (d, y - 0.12), ha="center", va="center", fontsize=7, color="white",
                            fontweight="bold", bbox=dict(boxstyle="round,pad=0.25", fc=col, ec=col),
                            arrowprops=dict(arrowstyle="-", color=col), zorder=6)
    # 주공정 연결 적색 점선
    A = {a.id: a for a in res.acts}
    yprep = 1.5
    first = [a for a in res.acts if a.critical and a.work > 0 and a.es == res.work_start]
    for a in first:
        if a.id in ypos:
            ax.plot([a.es, a.es], [yprep, ypos[a.id]], ls="--", color=RED, lw=1.3, zorder=2)
    for a in res.acts:
        if not (a.critical and a.id in ypos):
            continue
        for r in a.preds:
            pa = A[r.pred]
            if pa.critical and pa.id in ypos:
                ax.plot([a.es, a.es], [ypos[pa.id], ypos[a.id]], ls="--", color=RED, lw=1.3, zorder=2)
    last = [a for a in res.acts if a.critical and a.ef == res.finish_day and a.id in ypos]
    for a in last:
        ax.plot([a.ef + 1, a.ef + 1], [ypos[a.id], n - 0.5], ls="--", color=RED, lw=1.3, zorder=2)
    # 공정률
    yp = n + 0.5
    ax.text(x0 - (x1 - x0) * 0.005, yp + 0.5, "공정률(%)", ha="right", va="center", fontsize=8.5, fontweight="bold", clip_on=False)
    ax.axhline(n, color="#555555", lw=1)
    ax.axhline(n + 1, color="#E0E0E0", lw=0.5)
    for (y_, m_, pct, cum) in prog:
        xc = mday(y_, m_) + calendar.monthrange(y_, m_)[1] / 2
        ax.text(xc, yp, f"{pct:.0f}%", ha="center", va="center", fontsize=7)
        ax.text(xc, yp + 1, f"{cum:.0f}%", ha="center", va="center", fontsize=7, color=NAVY)
    # 머리 정보
    t = title or f"{p.name} 예정공정표"
    fig.text(0.5, 1 - 0.18 / fig_h, t, ha="center", va="top", fontsize=15, fontweight="bold")
    info = [f"■ 공사예정공정 ( {start.year}년 {start.month}월 ~ {end.year}년 {end.month}월 )",
            f"■ 착공일로부터 {total / R.DAYS_PER_MONTH:.1f}개월 ({total:,}일)",
            f"■ 위치 : {p.g('사업위치', '-')}"]
    for k, s in enumerate(info):
        fig.text(0.04, 1 - (0.62 + k * 0.2) / fig_h, s, fontsize=9.5, va="top")
    lx, ly = 0.70, 1 - 0.62 / fig_h
    fig.text(lx, ly, "범례", fontsize=9, fontweight="bold", va="top")
    lg = fig.add_axes([lx + 0.04, ly - 0.62 / fig_h, 0.2, 0.62 / fig_h])
    lg.set_xlim(0, 10)
    lg.set_ylim(3.2, 0)
    lg.axis("off")
    lg.add_patch(Rectangle((0.2, 0.35), 3, 0.5, color=GREEN))
    lg.add_patch(FancyArrowPatch((0.2, 0.6), (3.2, 0.6), arrowstyle="-|>", mutation_scale=8, color=RED, lw=1.5))
    lg.text(3.6, 0.6, "주공정선 (Critical Path)", va="center", fontsize=8)
    lg.add_patch(FancyArrowPatch((0.2, 1.6), (3.2, 1.6), arrowstyle="-|>", mutation_scale=8, color="black", lw=1.1))
    lg.text(3.6, 1.6, "비주공정", va="center", fontsize=8)
    lg.add_patch(Rectangle((0.2, 2.4), 3, 0.4, color=GREY))
    lg.text(3.6, 2.6, "총여유(TF, Total Float)", va="center", fontsize=8)
    fig.savefig(path_png, dpi=170)
    if path_pdf:
        fig.savefig(path_pdf)
    plt.close(fig)


# ── Excel 공정표 ──────────────────────────────────────────────────────────────
def write_sheet(ws, res, p):
    """10일(상·중·하순) 단위 막대 공정표를 시트에 작성."""
    cal = res.cal
    total = res.total_days
    start, end = cal.day_to_date(0), cal.day_to_date(total - 1)
    months = list(month_iter(start, end))
    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    F = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
    fills = {"crit": F("E53935"), "non": F("1E88E5"), "tf": F("D0D0D0"), "prep": F("FFB74D"), "hdr": F("1F4E78"), "grp": F("DDEBF7")}
    ws["A1"] = f"{p.name} — 주공정(Critical Path) 공정표"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"공사기간 {total:,}일(≒{total / R.DAYS_PER_MONTH:.1f}개월)  {start:%Y.%m.%d} ~ {end:%Y.%m.%d}   "
                "범례: 적색=주공정, 청색=비주공정, 회색=총여유(TF), 주황=준비·정리")
    hdr = ["ID", "구분", "공정명", "작업일수", "착수", "완료", "기간(일)", "TF", "CP"]
    nfix = len(hdr)
    slots = []  # (y, m, k, d_from, d_to)
    for y, m in months:
        nd = calendar.monthrange(y, m)[1]
        for k, (a, b) in enumerate(((1, 10), (11, 20), (21, nd))):
            slots.append((y, m, k, cal.date_to_day(dt.date(y, m, a)), cal.date_to_day(dt.date(y, m, b))))
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h)
        ws.merge_cells(start_row=4, start_column=j, end_row=5, end_column=j)
    col = nfix + 1
    for y, m in months:
        ws.cell(4, col, f"{str(y)[2:]}.{m}")
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 2)
        for k, s in enumerate(("상", "중", "하")):
            ws.cell(5, col + k, s)
        col += 3
    for r in (4, 5):
        for j in range(1, col):
            c = ws.cell(r, j)
            c.fill, c.font, c.border = fills["hdr"], Font(color="FFFFFF", bold=True, size=8), box
            c.alignment = Alignment(horizontal="center", vertical="center")
    widths = [6, 7, 28, 8, 10, 10, 7, 5, 4]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for j in range(nfix + 1, col):
        ws.column_dimensions[get_column_letter(j)].width = 2.6
    r = 6

    def paint(row, a_day, b_day, fill):
        for i, (_, _, _, s0, s1) in enumerate(slots):
            if s1 >= a_day and s0 <= b_day:
                ws.cell(row, nfix + 1 + i).fill = fill

    def line(vals, fill_row=None):
        nonlocal r
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.border = box
            c.font = Font(size=9, bold=fill_row is not None)
            if fill_row is not None:
                c.fill = fill_row
        for j in range(nfix + 1, col):
            ws.cell(r, j).border = box
        r += 1
        return r - 1

    rr = line(["", "준비", "준비기간", "", f"{start:%Y-%m-%d}", f"{cal.day_to_date(res.prep_end):%Y-%m-%d}", res.prep, "", "●"])
    paint(rr, 0, res.prep_end, fills["prep"])
    for kind, label, obj in chart_rows(res):
        if kind == "field":
            line(["", label.replace("■ ", ""), "", "", "", "", "", "", ""], fills["grp"])
        elif kind == "act":
            a = obj
            if a.work <= 0:
                continue
            rr = line([a.id, a.field, a.name, round(a.work), f"{cal.day_to_date(a.es):%Y-%m-%d}",
                       f"{cal.day_to_date(a.ef):%Y-%m-%d}", a.cal_days, a.tf, "●" if a.critical else ""])
            paint(rr, a.es, a.ef, fills["crit"] if a.critical else fills["non"])
            if a.tf > 0:
                paint(rr, a.ef + 1, a.lf, fills["tf"])
    rr = line(["", "정리", "정리기간", "", f"{cal.day_to_date(res.finish_day + 1):%Y-%m-%d}", f"{end:%Y-%m-%d}", res.cleanup, "", "●"])
    paint(rr, res.finish_day + 1, total - 1, fills["prep"])
    prog = progress_curve(res)
    r1 = line(["", "공정률", "월간 계획 공정률(%)", "", "", "", "", "", ""], fills["grp"])
    r2 = line(["", "공정률", "누계 공정률(%)", "", "", "", "", "", ""], fills["grp"])
    for i, (y, m, pct, cum) in enumerate(prog):
        c0 = nfix + 1 + i * 3
        ws.merge_cells(start_row=r1, start_column=c0, end_row=r1, end_column=c0 + 2)
        ws.merge_cells(start_row=r2, start_column=c0, end_row=r2, end_column=c0 + 2)
        ws.cell(r1, c0, round(pct, 1)).font = Font(size=8)
        ws.cell(r2, c0, round(cum, 1)).font = Font(size=8, bold=True)
    ws.freeze_panes = ws.cell(6, nfix + 1)
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
