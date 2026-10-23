"""공사기간 적정성 검토 프로그램 CLI

  python3 -m cpa template <입력서식.xlsx>             빈 입력서식 작성
  python3 -m cpa run <입력.xlsx> -o <출력폴더> [--no-pdf]  검토 실행 → 보고서·부록·공정표·검토워크북
  python3 -m cpa fill <입력.xlsx> --items 수량.csv --acts 공정.csv [--cases 사례.csv]  CSV 일괄입력
  python3 -m cpa composite 공사A=a.xlsx 공사B=b.xlsx@2027-06-01 [--name 복합공사명] -o <출력폴더>   복합공사
  python3 -m cpa qty --work 처리시설공사 토공.xlsx 구조물공.xlsx --work 관공사 관로.xlsx --start 2027-03-01 --end 2028-10-31 -o <출력>
  python3 -m cpa web [--port 8000]                    웹 화면 실행(파일 추가·검토·결과·보고서·출력)
  python3 -m cpa db [검색어]                          부록4 1일작업량 DB 검색
  python3 -m cpa stations [검색어]                     부록3 기상관측지점 검색
"""
import argparse
import logging
import os
import sys

logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)


def cmd_run(a):
    from . import service
    from .summary import fmt_date

    os.makedirs(a.out, exist_ok=True)
    res = service.calc(a.input)
    log = res.log
    log("[출력] 성과품 작성", "출력")
    service.chart(res, a.out)
    s = service.reports(res, a.out)
    if not a.no_pdf:
        service.export_pdf(a.out, log)
    log(f"[완료] 공사기간 {s['공사기간_일']:,}일(≒{s['공사기간_개월']}개월), {fmt_date(res.p.start_date)} ~ {s['준공']}, 적정성 '{s['적정성']}'", "완료")
    with open(os.path.join(a.out, "작업경과.log"), "w", encoding="utf-8") as f:
        for e in log.entries:
            f.write(f"{e['시각']} [{e['단계']}] {'[경고] ' if e['구분'] == '경고' else ''}{e['내용']}\n")
    print("\n성과품:", *sorted(x for x in os.listdir(a.out) if not x.startswith("_")), sep="\n  ")


def cmd_composite(a):
    """복합공사: 공사별 입력서식 → 통합 + 공사별 검토, 성과품 작성."""
    import datetime as dt

    from . import service
    from .combine import build
    from .pipeline import RunLog
    specs = []
    for tok in a.projects:                    # 공사명=파일[@착공일][+준비]
        nm, _, rest = tok.partition("=")
        own = rest.endswith("+준비")
        rest = rest[:-3] if own else rest
        path, _, start = rest.partition("@")
        specs.append({"name": nm, "path": path, "start": dt.date.fromisoformat(start) if start else None, "own_prep": own})
    os.makedirs(a.out, exist_ok=True)
    comb = os.path.join(a.out, "통합_복합공사_입력.xlsx")
    info = build(specs, comb, a.name)
    for w in info["warnings"]:
        print("[경고]", w)
    log = RunLog(echo=True)
    res = service.calc_composite(comb, [{"name": s["name"], "path": s["path"]} for s in specs], log)
    service.chart(res, a.out)
    service.reports(res, a.out)
    if not a.no_pdf:
        service.export_pdf(a.out, log)
    for x in service.sub_rows(res):
        print(f"  {x['공사명']}: 공사기간 {x['독립_공사기간_일']}일, A {x['A_일']} / B {x['B_일']} ({x['편차_%']}%) → {x['판정']}")
    print(f"통합 공사기간 {res.adopted_total:,}일, 성과품: {a.out}")


def cmd_qty(a):
    """수량첨부 검토: 공정별 수량파일 → 입력서식 자동 작성 → 검토·성과품."""
    import json

    from . import service, wizard
    from .qtyparse import category_from_filename
    st = wizard.default_state()
    for grp in a.work:
        name, files = grp[0], grp[1:]
        st["works"].append({"name": name, "files": [{"file": os.path.abspath(f), "orig": os.path.basename(f),
                                                     "category": category_from_filename(f), "sheets_off": []} for f in files]})
    ov = {"사업명": a.name, "공사명": a.name, "착공예정일": a.start, "준공예정일": a.end, "총공사비(백만원)": a.cost,
          "사업위치": a.loc, "발주청": a.org}
    if a.ov:
        with open(a.ov, encoding="utf-8") as f:
            ov.update(json.load(f))
    st["ov"] = {k: v for k, v in ov.items() if v not in (None, "")}
    os.makedirs(a.out, exist_ok=True)
    inp = os.path.join(a.out, "검토입력.xlsx")
    info = wizard.build_input(st, "", inp)
    print(f"수량 {len(info['rows'])}행 → 공정 {len(info['acts'])}개, 미반영 {len(info['excluded'])}행(검토입력.xlsx 6_제외수량) → {inp}")
    res = service.calc(inp)
    service.chart(res, a.out)
    s = service.reports(res, a.out)
    if not a.no_pdf:
        service.export_pdf(a.out, res.log)
    print(f"공사기간 {s['공사기간_일']:,}일(≒{s['공사기간_개월']}개월), 적정성 '{s['적정성']}' — 성과품: {a.out}")


def cmd_web(a):
    import socket
    import threading
    import time
    import webbrowser

    import uvicorn

    from .web.app import app
    host = "127.0.0.1" if a.host in ("localhost",) else a.host
    url = f"http://{'127.0.0.1' if host in ('0.0.0.0', '127.0.0.1') else host}:{a.port}"
    with socket.socket() as s:
        if s.connect_ex(("127.0.0.1", a.port)) == 0:
            print(f"[안내] {a.port}번 포트가 이미 사용 중입니다. 이미 실행 중이면 브라우저에서 {url} 로 접속하거나, --port 8001 처럼 다른 포트로 실행하십시오.")
            if a.open:
                webbrowser.open(url)
            return

    def opener():
        for _ in range(120):                     # 서버가 준비될 때까지 최대 60초 대기
            with socket.socket() as s:
                if s.connect_ex(("127.0.0.1", a.port)) == 0:
                    webbrowser.open(url)
                    return
            time.sleep(0.5)

    if a.open:
        threading.Thread(target=opener, daemon=True).start()
    print(f"공사기간 적정성 검토 웹 실행 중: {url}  (종료: Ctrl+C)", flush=True)
    uvicorn.run(app, host=host, port=a.port, log_level="warning")


def cmd_template(a):
    from .project import write_template
    write_template(a.path)
    print("입력서식 작성:", a.path)


def cmd_fill(a):
    """CSV(머리글=입력서식 열 이름)를 입력서식의 3_수량/4_공정/5_유사사례 시트에 기록(기존 행 대체)."""
    import csv
    import openpyxl
    from .project import ACT_COLS, CASE_COLS, ITEM_COLS
    wb = openpyxl.load_workbook(a.input)
    for csvp, sheet, cols in ((a.items, "3_수량", ITEM_COLS), (a.acts, "4_공정", ACT_COLS), (a.cases, "5_유사사례", CASE_COLS)):
        if not csvp:
            continue
        with open(csvp, encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
        ws = wb[sheet]
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.value = None
        for i, row in enumerate(rows, 2):
            for j, k in enumerate(cols, 1):
                v = row.get(k, "")
                if v not in ("", None):
                    try:
                        v = float(v) if any(ch in v for ch in ".eE") else int(v)
                    except ValueError:
                        pass
                    ws.cell(i, j, v)
        print(f"{sheet}: {len(rows)}행 기록")
    wb.save(a.input)


def cmd_db(a):
    from . import refdata as R
    from .workdays import dice
    q = " ".join(a.query or [])
    items = R.PRODUCTIVITY["items"]
    if q:
        items = sorted(items, key=lambda it: -dice(q, it["group"] + it["name"] + it["condition"]))[:25]
    for it in items:
        print(f"{it['no']:>4} | {it['field']:<4} | {it['group']:<8} | {it['name']:<26} | {it['condition']:<28} | {it['display']:<14} | {it['source']}")


def cmd_stations(a):
    from . import refdata as R
    q = " ".join(a.query or [])
    for k, v in sorted(R.WEATHER["stations"].items(), key=lambda kv: (kv[1]["region"], kv[0])):
        if not q or q in k or q in v["region"]:
            print(f"{k:<8} {v['region']:<8} 지점번호 {v['code']}")


def main(argv=None):
    ap = argparse.ArgumentParser(prog="cpa", description="공사기간 적정성 검토 (2026 적정 공사기간 확보를 위한 가이드라인)")
    sp = ap.add_subparsers(dest="cmd", required=True)
    r = sp.add_parser("run")
    r.add_argument("input")
    r.add_argument("-o", "--out", default="output")
    r.add_argument("--no-pdf", action="store_true")
    r.set_defaults(f=cmd_run)
    cp = sp.add_parser("composite", help="복합공사: 공사명=입력.xlsx[@착공일][+준비] 를 2개 이상")
    cp.add_argument("projects", nargs="+")
    cp.add_argument("--name", default="복합공사")
    cp.add_argument("-o", "--out", default="output/복합공사")
    cp.add_argument("--no-pdf", action="store_true")
    cp.set_defaults(f=cmd_composite)
    q = sp.add_parser("qty", help="수량첨부 검토: --work 공정명 수량1.xlsx 수량2.xlsx ... (공정: 처리시설공사·구내배관공사·관공사·배수설비 공사·펌프장 공사)")
    q.add_argument("--work", nargs="+", action="append", required=True, metavar="공정명 파일")
    q.add_argument("--name", default="수량첨부 검토")
    q.add_argument("--start", required=True, help="착공예정일 YYYY-MM-DD")
    q.add_argument("--end", help="준공예정일 YYYY-MM-DD (목표기간)")
    q.add_argument("--cost", type=float, help="총공사비(백만원)")
    q.add_argument("--loc", help="사업위치(관측지점 자동선택)")
    q.add_argument("--org", help="발주청")
    q.add_argument("--ov", help="공사개요 추가값 JSON (1_공사개요 항목명: 값)")
    q.add_argument("-o", "--out", default="output/수량첨부")
    q.add_argument("--no-pdf", action="store_true")
    q.set_defaults(f=cmd_qty)
    wv = sp.add_parser("web")
    wv.add_argument("--host", default="127.0.0.1")
    wv.add_argument("--port", type=int, default=8000)
    wv.add_argument("--open", action="store_true", help="서버 준비 후 브라우저 자동 실행")
    wv.set_defaults(f=cmd_web)
    t = sp.add_parser("template")
    t.add_argument("path")
    t.set_defaults(f=cmd_template)
    fl = sp.add_parser("fill")
    fl.add_argument("input")
    fl.add_argument("--items")
    fl.add_argument("--acts")
    fl.add_argument("--cases")
    fl.set_defaults(f=cmd_fill)
    d = sp.add_parser("db")
    d.add_argument("query", nargs="*")
    d.set_defaults(f=cmd_db)
    s = sp.add_parser("stations")
    s.add_argument("query", nargs="*")
    s.set_defaults(f=cmd_stations)
    a = ap.parse_args(argv)
    a.f(a)


if __name__ == "__main__":
    sys.exit(main())
