"""검토과정 워크북 — 참고 검토엑셀의 시트 순서(①→⑦)대로 작업경과와 결과를 기록한다."""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import gantt
from . import refdata as R
from .summary import fmt_date, monthly_rows

THIN = Side(style="thin", color="AAAAAA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HF = PatternFill("solid", fgColor="1F4E78")
SF = PatternFill("solid", fgColor="DDEBF7")
RF = PatternFill("solid", fgColor="E2EFDA")
WF = PatternFill("solid", fgColor="FCE4D6")
CF = PatternFill("solid", fgColor="FFE0E0")


class W:
    def __init__(self, ws):
        self.ws, self.r = ws, 1

    def title(self, t, sub=None):
        self.ws.cell(self.r, 1, t).font = Font(bold=True, size=14, color="1F4E78")
        self.r += 1
        if sub:
            self.ws.cell(self.r, 1, sub).font = Font(size=9, color="555555")
            self.r += 1
        self.r += 1

    def section(self, t):
        c = self.ws.cell(self.r, 1, t)
        c.font = Font(bold=True, size=11)
        self.r += 1

    def table(self, header, rows, widths=None, fills=None, fmt=None):
        for j, h in enumerate(header, 1):
            c = self.ws.cell(self.r, j, h)
            c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=9), HF, BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.ws.row_dimensions[self.r].height = 30
        self.r += 1
        for i, row in enumerate(rows):
            for j, v in enumerate(row, 1):
                c = self.ws.cell(self.r, j, v)
                c.border, c.font = BOX, Font(size=9)
                c.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and len(v) > 25)
                if fmt and j - 1 in fmt and isinstance(v, (int, float)):
                    c.number_format = fmt[j - 1]
                if fills and fills[i]:
                    c.fill = fills[i]
            self.r += 1
        if widths:
            for j, w in enumerate(widths, 1):
                cur = self.ws.column_dimensions[get_column_letter(j)].width or 0
                self.ws.column_dimensions[get_column_letter(j)].width = max(cur, w)
        self.r += 1

    def kv(self, pairs):
        for k, v in pairs:
            a = self.ws.cell(self.r, 1, k)
            a.font, a.fill, a.border = Font(bold=True, size=9), SF, BOX
            b = self.ws.cell(self.r, 2, v)
            b.border, b.font = BOX, Font(size=9)
            b.alignment = Alignment(wrap_text=True, vertical="top")
            self.r += 1
        self.ws.column_dimensions["A"].width = max(self.ws.column_dimensions["A"].width or 0, 26)
        self.ws.column_dimensions["B"].width = max(self.ws.column_dimensions["B"].width or 0, 60)
        self.r += 1

    def note(self, t):
        self.ws.cell(self.r, 1, t).font = Font(size=9, italic=True, color="555555")
        self.r += 1


def write(res, path):
    p, fin = res.p, res.final
    cal = fin.cal
    wb = openpyxl.Workbook()

    # 0 작업경과
    w = W(wb.active)
    wb.active.title = "0_작업경과"
    w.title("작업경과 로그", f"{p.name} — 검토 실행 순서(①→⑦)별 처리내용과 결과")
    w.table(["시각", "단계", "구분", "내용"], [[e["시각"], e["단계"], e["구분"], e["내용"]] for e in res.log.entries],
            [9, 18, 7, 140], [WF if e["구분"] == "경고" else None for e in res.log.entries])

    # ① 공사개요
    w = W(wb.create_sheet("①공사개요"))
    w.title("① 공사 개요 및 기본조건", R.GUIDELINE)
    w.section("[ 입력 조건 ]")
    w.kv([(k, (fmt_date(v) if hasattr(v, "strftime") else v)) for k, v in p.ov.items() if v not in (None, "")])
    w.section("[ 산정 결과 요약 ]")
    total = res.adopted_total
    w.kv([("① 준비기간", f"{fin.prep}일"), ("② 비작업일수(CP)", f"{fin.cp_nonwork:.0f}일"),
          ("③ 작업일수(CP)", f"{fin.cp_work:.0f}일"), ("④ 정리기간", f"{fin.cleanup}일"),
          ("공사여건 보정", f"{res.correction_days}일"),
          ("▶ 총 공사기간", f"{total:,}일 ≒ {total / R.DAYS_PER_MONTH:.1f}개월"),
          ("착공 ~ 준공", f"{fmt_date(cal.day_to_date(0))} ~ {fmt_date(cal.day_to_date(total - 1))}"),
          ("적정성 판정", f"{res.adequacy['verdict']} (비교기준: {res.adequacy['basis']})")])

    # ②-1 기후비작업근거
    w = W(wb.create_sheet("②-1 기후비작업근거"))
    st = res.setup.station_data()
    w.title("②-1 기후 비작업일수 산정근거", f"부록3 기상조건별·지역별 비작업일수(2015~2024) — 관측지점 '{res.setup.station}' ({st['region']}, 지점번호 {st['code']})")
    w.section("[ 관측지점 13개 기상조건 월평균 비작업일수 + 미세먼지(입력) ]")
    rows = []
    for c in range(1, 15):
        vals = [res.setup.cond_value(c, m) for m in range(1, 13)]
        rows.append([c, R.COND_GROUP[c], R.CONDITIONS[c]] + vals + [round(sum(vals), 1)])
    w.table(["번호", "구분", "기상조건"] + [f"{m}월" for m in range(1, 13)] + ["계"], rows, [6, 9, 44] + [6] * 13)
    w.section("[ 기상프로필별 적용 기상조건·적용률 ]")
    pn = list(res.setup.profiles)
    w.table(["번호", "기상조건"] + pn, [[c, R.CONDITIONS[c]] + [res.setup.profiles[x].get(c) or "" for x in pn] for c in range(1, 15)])
    for x in pn:
        ann = res.setup.annual_climate(x)
        if not ann:
            continue
        w.section(f"[ 프로필 '{x}' 월별 기상 비작업일수(A) — {R.PROFILE_DESC.get(x, '')} ]")
        rr = [[R.COND_SHORT[c]] + v + [round(sum(v), 1)] for c, v in ann.items()]
        tot = [round(sum(v[m] for v in ann.values()), 1) for m in range(12)]
        rr.append(["합계"] + tot + [round(sum(tot), 1)])
        w.table(["기상조건"] + [f"{m}월" for m in range(1, 13)] + ["계"], rr, fills=[None] * (len(rr) - 1) + [RF])

    # ② 비작업일수
    w = W(wb.create_sheet("②비작업일수"))
    w.title("② 월별 비작업일수 산정", "비작업일수 = A + B − C,  C = A×B÷달력일수(반올림),  주40시간 하한 월 8일. 공사진행일수 비례로 공사기간 반영")
    mr = monthly_rows(fin, res.setup)
    hdr = ["년", "월", "구분", "달력일수", "공사진행일", "지배프로필(CP)", "기후(A)", "공휴일(B)", "중복(C)", "A+B−C", "적용비작업", "작업가능일", "가동률", "공사기간 반영 비작업", "반영 작업일"]
    w.table(hdr, [[r["년"], r["월"], r["구분"], r["달력일수"], r["공사일수"], r["지배프로필"], r["A"], r["B"], r["C"], r["소계"], r["적용"],
                   r["작업가능"], r["가동률"], r["반영비작업"], r["반영작업"]] for r in mr],
            [6, 4, 10, 8, 8, 12, 7, 7, 7, 7, 8, 8, 7, 10, 9], fmt={12: "0.0%"})
    w.note(f"※ 주공정 비작업일수 합계 {sum(r['반영비작업'] for r in mr):.1f}일, 작업일수 {sum(r['반영작업'] for r in mr):.1f}일 (준비·정리기간은 비작업일수 미계상)")
    for x in pn:
        w.section(f"[ 프로필 '{x}' 월별 산정표 ]")
        rr = []
        for r in mr:
            mi = res.setup.month(x, r["년"], r["월"])
            rr.append([r["년"], r["월"], mi.days, mi.A, mi.B, mi.C, mi.raw, mi.applied, mi.workable, mi.rate])
        w.table(["년", "월", "달력일수", "A", "B", "C", "A+B−C", "적용", "작업가능", "가동률"], rr, fmt={9: "0.0%"})

    # ③ 1일작업량 매칭
    w = W(wb.create_sheet("③1일작업량매칭"))
    w.title("③ 1일 작업량 결정 (부록4 DB 매칭)", "우선순위: 수동입력(품셈·실적) > DB지정 > 자동매칭(규칙/유사도). 자동매칭 항목은 검토자가 반드시 확인")
    items = [it for a in fin.acts for it in res.final_items.get(a.id, [])]
    rows, fl = [], []
    for it in items:
        db = it.db or {}
        rows.append([it.no, it.act, it.name, it.spec, it.unit, it.match_note, it.confidence, db.get("no"), db.get("group"),
                     db.get("name"), db.get("condition"), db.get("display"), it.base_rate, it.source, it.warn])
        fl.append(WF if it.warn else (CF if str(it.match_note).startswith("유사도") else None))
    w.table(["No", "공정", "공종명", "규격", "단위", "결정방법", "신뢰도", "DB번호", "DB대공정", "DB공종", "DB작업조건", "부록4 표기",
             "1일작업량", "출처", "검토사항"], rows, [5, 6, 26, 22, 6, 14, 6, 6, 10, 22, 20, 12, 9, 18, 30], fl)
    w.note("※ 분홍색 = 유사도 매칭(확인 필요), 주황색 = 경고")

    # ④ 작업일수산정
    w = W(wb.create_sheet("④작업일수산정"))
    w.title("④ 작업일수 산정", "작업일수 = ROUNDUP(수량 ÷ (1일작업량 × 보정계수 × 작업조)), 1일 미만은 1일 / 보정계수 = 현장여건×공사규모×지질×기상기후×개별")
    hdr = ["No", "공정", "공종명", "규격", "단위", "수량", "1일작업량", "현장", "규모", "지질", "기상", "보정계수", "적용1일작업량",
           "작업조", "순작업일", "간접", "보정", "합산작업일", "산식", "비고"]
    for a in fin.acts:
        its = res.final_items.get(a.id, [])
        w.section(f"[{a.id}] {a.field} · {a.group} · {a.name}  — 공정 작업일수 {a.work:.0f}일 ({a.profile})"
                  + (f", 직접입력 {a.direct_days:g}일" if a.direct_days else ""))
        if its:
            w.table(hdr, [[it.no, it.act, it.name, it.spec, it.unit, it.qty, it.base_rate, *it.factors, it.factor,
                           round(it.rate, 3) if it.rate else None, it.crews, it.days, it.indirect, it.correction,
                           "중복" if it.parallel else it.total, it.formula, it.note] for it in its],
                    [5, 6, 26, 20, 6, 9, 9, 5, 5, 5, 5, 7, 9, 6, 7, 5, 5, 8, 34, 24])

    # ⑤ CP구성
    w = W(wb.create_sheet("⑤CP구성"))
    w.title("⑤ 주공정(CP) 구성 — CPM", "ES/EF=최조개시/완료, LS/LF=최늦개시/완료, TF=총여유. 공정별 기상프로필 가동률 달력으로 달력기간 산정")
    rows, fl = [], []
    for a in sorted(fin.acts, key=lambda a: (a.es, a.ef)):
        d = cal.day_to_date
        rows.append([a.id, a.field, a.group, a.name, a.profile, a.work, ", ".join(r.label() for r in a.preds) or "착공(준비 후)",
                     fmt_date(d(a.es)), fmt_date(d(a.ef)), fmt_date(d(a.ls)), fmt_date(d(a.lf)), a.cal_days,
                     round(a.cal_days - a.work, 1) if a.work else 0, a.tf, "●" if a.critical else ""])
        fl.append(CF if a.critical else None)
    w.table(["ID", "분야", "대공정", "공정명", "기상프로필", "작업일수", "선행관계", "ES", "EF", "LS", "LF", "공종별 공사기간",
             "비작업일수", "TF", "CP"], rows, [6, 6, 10, 28, 10, 8, 26, 11, 11, 11, 11, 9, 9, 6, 4], fl)
    w.note(f"주공정: {' → '.join(fin.cp_ids)}")
    w.note("공종별 공사기간 = 작업일수 + 비작업일수 (가이드라인 제2장 3.(3))")
    w.section("[ 주공정 공종별 공사기간 ]")
    crit = [fin.act(i) for i in fin.cp_ids]
    w.table(["구분", "공정", "작업일수", "비작업일수", "공사기간"],
            [["준비기간", "", "", "", fin.prep]] + [[a.id, a.name, a.work, a.cal_days - a.work, a.cal_days] for a in crit]
            + [["정리기간", "", "", "", fin.cleanup], ["계(달력)", "중첩·연결 반영", round(fin.cp_work), round(fin.cp_nonwork), fin.total_days]])

    # ⑤-1 공기조정
    w = W(wb.create_sheet("⑤-1 공기조정"))
    w.title("⑤-1 공사기간 조정 (작업조 증대 · 동시수행 공정 중첩)")
    if res.target_days:
        w.kv([("목표기간", f"{res.target_days}일 ({res.target_desc})"), ("기본 산정", f"{res.base.total_days}일"),
              ("조정 결과", f"{fin.total_days}일" if res.compress_history else "조정 불필요"),
              ("목표 충족", "충족" if fin.total_days <= res.target_days else "미충족")])
        if res.compress_history:
            h = res.compress_history
            w.table(list(h[0].keys()), [list(x.values()) for x in h], [6, 8, 90, 10, 10, 8])
            chg = []
            base_items = {it.no: it for v in res.items_by_act.values() for it in v}
            for v in res.final_items.values():
                for it in v:
                    b = base_items.get(it.no)
                    if b and b.crews != it.crews:
                        chg.append([it.no, it.act, it.name, b.crews, it.crews, b.days, it.days])
            if chg:
                w.section("[ 작업조 변경 내역 ]")
                w.table(["No", "공정", "공종명", "당초 작업조", "조정 작업조", "당초 작업일", "조정 작업일"], chg)
            rel = []
            for a in fin.acts:
                b = res.base.act(a.id)
                if [r.label() for r in a.preds] != [r.label() for r in b.preds]:
                    rel.append([a.id, a.name, ", ".join(r.label() for r in b.preds), ", ".join(r.label() for r in a.preds)])
            if rel:
                w.section("[ 공정 중첩(선후관계) 변경 내역 ]")
                w.table(["ID", "공정명", "당초 선행관계", "조정 선행관계"], rel)
    else:
        w.note("목표기간 기준 '없음' — 공기조정 미수행")

    # CP공정표
    gantt.write_sheet(wb.create_sheet("CP공정표"), fin, p)

    # ⑥ 공사기간종합
    w = W(wb.create_sheet("⑥공사기간종합"))
    w.title("⑥ 공사기간 종합", "공사기간 = ① 준비기간 + ② 비작업일수 + ③ 작업일수 + ④ 정리기간  (+ 공사여건 보정)")
    t = res.adopted_total
    w.table(["구분", "일수", "개월", "기준"],
            [["① 준비기간", fin.prep, fin.prep / R.DAYS_PER_MONTH, "제7조(준비기간)"],
             ["② 비작업일수", round(fin.cp_nonwork), fin.cp_nonwork / R.DAYS_PER_MONTH, "제8~10조(비작업일수)"],
             ["③ 작업일수", round(fin.cp_work), fin.cp_work / R.DAYS_PER_MONTH, "제11조(작업일수)"],
             ["④ 정리기간", fin.cleanup, fin.cleanup / R.DAYS_PER_MONTH, "제12조(정리기간)"],
             ["동절기 공사중지", round(fin.winter_stop_days), fin.winter_stop_days / R.DAYS_PER_MONTH, "준비·정리기간 포함"],
             ["공사여건 보정", res.correction_days, res.correction_days / R.DAYS_PER_MONTH, "제2장 2.(6)"],
             ["합계", t, t / R.DAYS_PER_MONTH, f"{fmt_date(cal.day_to_date(0))} ~ {fmt_date(cal.day_to_date(t - 1))}"]],
            [18, 10, 10, 40], [None] * 6 + [RF], fmt={2: "0.0"})
    if fin.winter_cal_days:
        w.note(f"※ 공사기간 내 동절기 달력일수 {fin.winter_cal_days}일 → 동절기 제외 공사기간 {t - fin.winter_cal_days}일 (적정성 비교용 A)")
    from .summary import progress_curve
    w.section("[ 월별 계획 공정률 ]")
    w.table(["년", "월", "월간(%)", "누계(%)"], [[y, m, round(a, 1), round(b, 1)] for y, m, a, b in progress_curve(fin)])

    # ⑦ 적정성검토
    w = W(wb.create_sheet("⑦적정성검토"))
    adq = res.adequacy
    w.title("⑦ 공사기간 적정성 검토", f"공사기간(A) vs 공사기간(B): |A−B|/B ≤ {adq['tol'] * 100:.0f}% 적정 — 가이드라인 제3장 2.(3)")
    rows, fl = [], []
    for r in adq["rows"]:
        rows.append([r["구분"], r["기준"], r["산정식"], r.get("변수", ""), round(r["Y"], 1) if r.get("Y") else "",
                     round(r["B"]) if r.get("B") else "", round(r["B"] / R.DAYS_PER_MONTH, 1) if r.get("B") else "",
                     adq["A"], f"{r['편차'] * 100:+.1f}%" if r.get("편차") is not None else "", r.get("판정", ""), r.get("비고", "")])
        fl.append(RF if r.get("주기준") else None)
    w.table(["구분", "기준", "산정식", "변수", "Y(일)", "B(일)", "B(개월)", "A(일)", "편차", "판정", "비고"], rows,
            [14, 36, 40, 30, 8, 8, 8, 8, 8, 14, 30], fl)
    w.kv([("종합판정", adq["verdict"]), ("주 비교기준", adq["basis"])])
    if p.cases:
        w.section("[ 최근 준공 동종·유사공사 ]")
        w.table(list(p.cases[0].keys()), [list(c.values()) for c in p.cases])

    from .service import sub_rows
    subs = sub_rows(res)
    if subs:
        w = W(wb.create_sheet("⑦-1 공사별 적정성"))
        w.title("⑦-1 복합공사 공사별 공사기간 및 적정성", "공사별 입력으로 독립 산정(준비·정리 포함), 통합 착수·완료는 통합 공정망 기준")
        w.table(list(subs[0].keys()), [list(x.values()) for x in subs])

    # 추가검토(동절기)
    w = W(wb.create_sheet("추가검토(동절기)"))
    w.title("추가검토 — 동절기 공사중지 반영 시 공사일정")
    if res.winter:
        wi = res.winter
        w.table(["구분", "공사기간(일)", "개월", "착공", "준공"],
                [["동절기 미반영(금회 산정)", fin.total_days, fin.months, fmt_date(cal.day_to_date(0)), fmt_date(fin.end_date)],
                 ["동절기 공사중지 반영", wi.total_days, wi.months, fmt_date(wi.cal.day_to_date(0)), fmt_date(wi.end_date)],
                 ["증감", wi.total_days - fin.total_days, wi.months - fin.months, "", ""]], fmt={2: "0.0"})
        w.note("※ 동절기 중지기간은 달력일수 전체를 비작업일수로 계상(기상조건 비작업일수와 중복계상 배제)")
    else:
        w.note(f"동절기 공사중지 설정 '{p.winter_mode}' — 추가검토 미수행")

    # 부록1
    w = W(wb.create_sheet("부록1 법정공휴일"))
    w.title("부록1 법정 공휴일수(2026-2035년)", R.HOLIDAYS["source"])
    w.table(["연도"] + [f"{m}월" for m in range(1, 13)] + ["계"], [[int(y)] + v + [sum(v)] for y, v in R.HOLIDAYS["data"].items()])
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
    wb.save(path)
