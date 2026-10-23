"""공사기간 적정성 검토 웹 — ① 파일 추가(단일/복합) ② 검토 실행 ③ 결과 확인 ④ 보고서 작성 ⑤ 보고서 출력

실행: python3 -m cpa web  →  http://localhost:8000
작업 데이터: <저장소>/web_data/<작업ID>/ (원본/·입력/·입력/공사별/·output/·meta.json)
"""
import datetime as dt
import json
import os
import shutil
import sys
import threading
import traceback
import uuid
from typing import List

import openpyxl
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse

from .. import refdata as R
from .. import service
from ..combine import build as combine_build
from ..pipeline import RunLog
from ..project import OVERVIEW
from . import wiz

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DATA = os.environ.get("CPA_WEB_DATA", os.path.join(ROOT, "web_data"))
STATIC = os.path.join(os.path.dirname(__file__), "static")
os.makedirs(DATA, exist_ok=True)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

app = FastAPI(title="공사기간 적정성 검토")
app.include_router(wiz.router)
LOCK = threading.Lock()
RUNTIME = {}   # job_id → {"log": RunLog, "res": Result, "busy": 단계명 or None, "error": str}

EDITABLE = ["공사명", "발주청", "사업위치", "총공사비(백만원)", "착공예정일", "시설유형(부록5)", "기상관측지점", "준비기간(일)",
            "정리기간(일)", "동절기 공사중지", "준비·정리기간 동절기 중지", "작업일수 산정단위(일)", "현장여건", "공사규모",
            "지질조건", "기상기후", "목표기간 기준", "발주처 요구 공사기간(개월)", "허용범위(±)", "적정성 판정기준"]
# 복합공사에서 공사별로 다르게 유지하는 항목(나머지는 공통조건으로 모든 공사에 적용)
PER_PROJECT = {"공사명", "사업위치", "총공사비(백만원)", "착공예정일", "시설유형(부록5)", "준비기간(일)", "정리기간(일)"}
OPTS = {k: o for k, _, _, o in OVERVIEW if o}
DESC = {k: d for k, _, d, _ in OVERVIEW if d}


# ── 저장소 ────────────────────────────────────────────────────────────────────
def jdir(jid):
    d = os.path.join(DATA, os.path.basename(jid))
    if not os.path.isdir(d):
        raise HTTPException(404, "작업을 찾을 수 없습니다.")
    return d


def meta(jid):
    with open(os.path.join(jdir(jid), "meta.json"), encoding="utf-8") as f:
        return json.load(f)


def save_meta(jid, m):
    with open(os.path.join(jdir(jid), "meta.json"), "w", encoding="utf-8") as f:
        json.dump(m, f, ensure_ascii=False, indent=1, default=str)


def rt(jid):
    with LOCK:
        return RUNTIME.setdefault(jid, {"log": None, "res": None, "busy": None, "error": None})


def ipath(jid, rel):
    return os.path.join(jdir(jid), "입력", rel)


def cur_input(jid):
    return ipath(jid, meta(jid)["current"])


def is_composite_run(m):
    return m.get("type") == "복합공사" and m["current"] == m.get("combined")


def outdir(jid):
    m = meta(jid)
    o = os.path.join(jdir(jid), "output", os.path.splitext(os.path.basename(m["current"]))[0])
    os.makedirs(o, exist_ok=True)
    return o


def _v(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else (v.isoformat() if isinstance(v, dt.date) else v)


def read_overview(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ov = {}
    for row in wb["1_공사개요"].iter_rows(min_row=2, values_only=True):
        if row[0] and not str(row[0]).startswith("■"):
            ov[str(row[0]).strip()] = _v(row[1])
    n_items = sum(1 for r in wb["3_수량"].iter_rows(min_row=2, values_only=True) if len(r) > 1 and r[1])
    n_acts = sum(1 for r in wb["4_공정"].iter_rows(min_row=2, values_only=True) if r and r[0])
    return ov, n_items, n_acts


def write_overview(path, values):
    wb = openpyxl.load_workbook(path)
    ws = wb["1_공사개요"]
    changed = []
    for row in ws.iter_rows(min_row=2):
        k = row[0].value
        if k not in values:
            continue
        v = values[k]
        if v in ("", None):
            v = None
        elif k == "착공예정일":
            v = dt.datetime.fromisoformat(str(v)[:10])
        else:
            sv = str(v)
            if sv.replace(".", "", 1).replace("-", "", 1).isdigit():
                v = float(sv)
                if v.is_integer() and k != "허용범위(±)":
                    v = int(v)
        if row[1].value != v:
            row[1].value = v
            if k == "착공예정일":
                row[1].number_format = "yyyy-mm-dd"
            changed.append(k)
    wb.save(path)
    return changed


def validate_input(path):
    names = openpyxl.load_workbook(path, read_only=True).sheetnames
    if "1_공사개요" in names and "4_공정" in names:
        return "입력서식"
    if "⑤CP구성" in names and "④-1 처리시설 작업일수" in names:
        return "검토엑셀"
    return None


def rebuild_combined(jid):
    """복합공사: 공사별 입력 → 통합 입력 재작성."""
    m = meta(jid)
    specs = [{"name": s["name"], "path": ipath(jid, s["file"]), "start": s.get("start") or None,
              "own_prep": bool(s.get("own_prep"))} for s in m["subs"]]
    info = combine_build(specs, ipath(jid, m["combined"]), m["name"])
    if m.get("combined_overrides"):
        write_overview(ipath(jid, m["combined"]), m["combined_overrides"])
    m["combine_warnings"] = info["warnings"]
    m["combine_subs"] = info["subs"]
    m["inputs"] = [m["combined"]] + [s["file"] for s in m["subs"]]
    if m["current"] not in m["inputs"]:
        m["current"] = m["combined"]
    save_meta(jid, m)
    rt(jid)["res"] = None
    return info


def new_job_dir():
    jid = dt.datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:4]
    d = os.path.join(DATA, jid)
    os.makedirs(os.path.join(d, "입력", "공사별"))
    os.makedirs(os.path.join(d, "원본"))
    return jid, d


async def save_upload(file: UploadFile, dest_dir):
    fn = os.path.basename(file.filename or "input.xlsx")
    if not fn.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, f"'{fn}': 엑셀(.xlsx) 파일만 업로드할 수 있습니다.")
    os.makedirs(dest_dir, exist_ok=True)
    p = os.path.join(dest_dir, fn)
    base, ext = os.path.splitext(p)
    k = 2
    while os.path.exists(p):
        p = f"{base}_{k}{ext}"
        k += 1
    with open(p, "wb") as f:
        f.write(await file.read())
    return p


# ── 백그라운드 작업 ───────────────────────────────────────────────────────────
def background(jid, step, fn):
    r = rt(jid)
    if r["busy"]:
        raise HTTPException(409, f"'{r['busy']}' 작업이 진행 중입니다.")
    r["busy"], r["error"] = step, None

    def work():
        try:
            fn(r)
        except SystemExit as e:
            r["error"] = str(e)
            r["log"] and r["log"].warn(str(e))
        except Exception as e:  # noqa: BLE001
            r["error"] = f"{type(e).__name__}: {e}"
            if r["log"]:
                r["log"].warn(r["error"])
            traceback.print_exc()
        finally:
            r["busy"] = None
            m = meta(jid)
            m.setdefault("history", []).append({"단계": step, "시각": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                                "결과": "오류" if r["error"] else "완료"})
            save_meta(jid, m)

    threading.Thread(target=work, daemon=True).start()
    return {"ok": True, "busy": step}


def do_calc(jid, log):
    m = meta(jid)
    if m.get("type") == "수량첨부":
        log("[0/9] 수량파일 ↔ 1일작업량 DB 매칭, 공정 구성 → 검토입력.xlsx 작성", "준비")
        info = wiz.build(jid)
        log(f"  수량 {len(info['rows'])}행(반영 {sum(1 for x in info['rows'] if x['use'])}행), 공정 {len(info['acts'])}개", "준비")
        if info["excluded"]:
            log.warn(f"작업일수 미반영 수량 {len(info['excluded'])}행(1일작업량 미매칭·검토자 제외) — 검토입력.xlsx '6_제외수량' 시트 확인, 2.3에서 DB지정·수동입력 가능")
    if is_composite_run(m):
        subs = [{"name": s["name"], "path": ipath(jid, s["file"])} for s in m["subs"]]
        return service.calc_composite(ipath(jid, m["combined"]), subs, log)
    return service.calc(cur_input(jid), log)


def ensure_res(jid, r):
    if r["res"] is None:
        r["log"] = RunLog(echo=False)
        r["log"]("이전 검토결과가 없어 검토를 다시 실행합니다.", "준비")
        r["res"] = do_calc(jid, r["log"])
    return r["res"]


# ── 페이지·정보 ───────────────────────────────────────────────────────────────
@app.get("/", response_class=HTMLResponse)
def index():
    with open(os.path.join(STATIC, "index.html"), encoding="utf-8") as f:
        return f.read()


@app.get("/classic", response_class=HTMLResponse)
def classic():
    """입력서식(1_공사개요~5_유사사례)·검토엑셀 업로드 방식(이전 화면)."""
    with open(os.path.join(STATIC, "classic.html"), encoding="utf-8") as f:
        return f.read()


@app.get("/api/info")
def info():
    return {"guideline": R.GUIDELINE, "stations": sorted(R.WEATHER["stations"]), "editable": EDITABLE,
            "per_project": sorted(PER_PROJECT), "options": {k: OPTS.get(k) for k in EDITABLE},
            "desc": {k: DESC.get(k, "") for k in EDITABLE}, "soffice": bool(service.soffice())}


@app.get("/api/template/{kind}")
def template(kind: str):
    p = {"blank": os.path.join(ROOT, "templates", "입력서식_템플릿.xlsx"),
         "example": os.path.join(ROOT, "examples", "예제_대죽공공폐수_입력.xlsx")}.get(kind)
    if not p:
        raise HTTPException(404)
    return FileResponse(p, filename=os.path.basename(p))


@app.post("/api/inspect")
async def inspect(file: UploadFile = File(...)):
    """첨부 전 확인: 파일 형식, 공사명·착공일·공정수·수량수 (복합공사 행 자동 채움·단일/복합 판별용)."""
    tmp = os.path.join(DATA, "_inspect")
    p = await save_upload(file, tmp)
    try:
        kind = validate_input(p)
        if kind != "입력서식":
            return {"kind": kind or "알수없음"}
        ov, ni, na = read_overview(p)
        return {"kind": kind, "공사명": ov.get("공사명"), "착공예정일": ov.get("착공예정일"), "공정수": na, "수량수": ni,
                "총공사비": ov.get("총공사비(백만원)"), "시설유형": ov.get("시설유형(부록5)")}
    finally:
        os.remove(p)


# ── 작업 생성 ─────────────────────────────────────────────────────────────────
@app.get("/api/jobs")
def jobs():
    out = []
    for jid in sorted(os.listdir(DATA), reverse=True):
        if jid.startswith("_"):
            continue
        try:
            m = meta(jid)
        except Exception:  # noqa: BLE001
            continue
        out.append({"id": jid, "name": m.get("name"), "created": m.get("created"), "kind": m.get("kind"),
                    "type": m.get("type", "단일공사"), "n_subs": len(m.get("subs") or []),
                    "busy": rt(jid)["busy"], "summary": m.get("summary")})
    return out


@app.post("/api/jobs")
async def create_single(file: UploadFile = File(...)):
    """단일공사(입력서식 1개) 또는 기존 검토엑셀(여러 사업 → 복합공사로 자동 인식)."""
    jid, d = new_job_dir()
    src = await save_upload(file, os.path.join(d, "원본"))
    fn = os.path.basename(src)
    kind = validate_input(src)
    if kind == "입력서식":
        shutil.copy(src, os.path.join(d, "입력", fn))
        ov, _, _ = read_overview(os.path.join(d, "입력", fn))
        save_meta(jid, {"name": ov.get("공사명") or fn, "file": fn, "type": "단일공사", "kind": "입력서식",
                        "inputs": [fn], "current": fn, "created": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "history": []})
        return {"id": jid, "type": "단일공사"}
    if kind == "검토엑셀":
        from tools.excel_to_input import main as convert  # noqa: PLC0415
        sub_dir = os.path.join(d, "입력", "공사별")
        convert(src, sub_dir)
        files = sorted(os.listdir(sub_dir))
        comb = next(f for f in files if "통합" in f)
        shutil.move(os.path.join(sub_dir, comb), os.path.join(d, "입력", comb))
        subs = []
        for f in files:
            if f == comb:
                continue
            ov, _, _ = read_overview(os.path.join(sub_dir, f))
            subs.append({"name": str(ov.get("공사명", f)).split(" (")[0], "file": f"공사별/{f}", "start": None, "own_prep": False})
        ov, _, _ = read_overview(os.path.join(d, "입력", comb))
        save_meta(jid, {"name": ov.get("공사명") or fn, "file": fn, "type": "복합공사", "kind": "검토엑셀 변환",
                        "subs": subs, "combined": comb, "inputs": [comb] + [s["file"] for s in subs], "current": comb,
                        "created": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "history": []})
        return {"id": jid, "type": "복합공사"}
    shutil.rmtree(d)
    raise HTTPException(400, "인식할 수 없는 엑셀입니다. 입력서식(1_공사개요·3_수량·4_공정 시트) 또는 검토엑셀(①~⑦ 시트)을 업로드하십시오.")


@app.post("/api/jobs/composite")
async def create_composite(name: str = Form(...), names: List[str] = Form(...), starts: List[str] = Form(None),
                           preps: List[str] = Form(None), files: List[UploadFile] = File(...)):
    """복합공사: 공사별 입력서식(수량·공정) 첨부 → 통합 입력 생성."""
    if len(files) < 2:
        raise HTTPException(400, "복합공사는 2개 이상의 공사 파일이 필요합니다.")
    if len(names) != len(files):
        raise HTTPException(400, "공사명과 파일 수가 일치하지 않습니다.")
    jid, d = new_job_dir()
    try:
        subs = []
        for k, (nm, f) in enumerate(zip(names, files)):
            p = await save_upload(f, os.path.join(d, "입력", "공사별"))
            if validate_input(p) != "입력서식":
                raise HTTPException(400, f"'{nm}' 첨부파일이 입력서식(1_공사개요·3_수량·4_공정)이 아닙니다.")
            subs.append({"name": nm.strip() or f"공사{k + 1}", "file": f"공사별/{os.path.basename(p)}",
                         "start": (starts[k] if starts and k < len(starts) and starts[k] else None),
                         "own_prep": bool(preps and k < len(preps) and preps[k] in ("Y", "true", "on", "1"))})
        if len({s["name"] for s in subs}) != len(subs):
            raise HTTPException(400, "공사명이 중복되었습니다.")
        save_meta(jid, {"name": name.strip() or "복합공사", "type": "복합공사", "kind": "공사별 첨부", "subs": subs,
                        "combined": "통합_복합공사.xlsx", "inputs": [], "current": "통합_복합공사.xlsx",
                        "created": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "history": []})
        info = rebuild_combined(jid)
    except HTTPException:
        shutil.rmtree(d)
        raise
    except Exception as e:  # noqa: BLE001
        shutil.rmtree(d)
        raise HTTPException(400, f"복합공사 구성 오류: {e}")
    return {"id": jid, "type": "복합공사", "warnings": info["warnings"]}


# ── 복합공사 구성 변경 ────────────────────────────────────────────────────────
def _need_comp(m):
    if m.get("type") != "복합공사" or m.get("kind") == "검토엑셀 변환":
        raise HTTPException(400, "공사별 첨부로 만든 복합공사에서만 구성을 변경할 수 있습니다.")


@app.post("/api/jobs/{jid}/subs")
async def add_sub(jid: str, name: str = Form(...), start: str = Form(""), own_prep: str = Form("N"),
                  file: UploadFile = File(...)):
    m = meta(jid)
    _need_comp(m)
    if name in [s["name"] for s in m["subs"]]:
        raise HTTPException(400, "같은 공사명이 이미 있습니다.")
    p = await save_upload(file, ipath(jid, "공사별"))
    if validate_input(p) != "입력서식":
        os.remove(p)
        raise HTTPException(400, "입력서식(1_공사개요·3_수량·4_공정) 파일이 아닙니다.")
    m["subs"].append({"name": name, "file": f"공사별/{os.path.basename(p)}", "start": start or None, "own_prep": own_prep == "Y"})
    save_meta(jid, m)
    return {"ok": True, "warnings": rebuild_combined(jid)["warnings"]}


@app.post("/api/jobs/{jid}/subs/{idx}/file")
async def replace_sub_file(jid: str, idx: int, file: UploadFile = File(...)):
    m = meta(jid)
    _need_comp(m)
    p = await save_upload(file, ipath(jid, "공사별"))
    if validate_input(p) != "입력서식":
        os.remove(p)
        raise HTTPException(400, "입력서식 파일이 아닙니다.")
    old = ipath(jid, m["subs"][idx]["file"])
    m["subs"][idx]["file"] = f"공사별/{os.path.basename(p)}"
    save_meta(jid, m)
    if os.path.exists(old) and os.path.abspath(old) != os.path.abspath(p):
        os.remove(old)
    return {"ok": True, "warnings": rebuild_combined(jid)["warnings"]}


@app.post("/api/jobs/{jid}/subs/{idx}")
def update_sub(jid: str, idx: int, body: dict):
    m = meta(jid)
    _need_comp(m)
    s = m["subs"][idx]
    if (body.get("name") or "").strip():
        if body["name"].strip() in [x["name"] for i, x in enumerate(m["subs"]) if i != idx]:
            raise HTTPException(400, "같은 공사명이 이미 있습니다.")
        s["name"] = body["name"].strip()
    if "start" in body:
        s["start"] = body["start"] or None
    if "own_prep" in body:
        s["own_prep"] = bool(body["own_prep"])
    save_meta(jid, m)
    return {"ok": True, "warnings": rebuild_combined(jid)["warnings"]}


@app.delete("/api/jobs/{jid}/subs/{idx}")
def delete_sub(jid: str, idx: int):
    m = meta(jid)
    _need_comp(m)
    if len(m["subs"]) <= 2:
        raise HTTPException(400, "복합공사는 최소 2개 공사가 필요합니다.")
    s = m["subs"].pop(idx)
    save_meta(jid, m)
    p = ipath(jid, s["file"])
    if os.path.exists(p):
        os.remove(p)
    return {"ok": True, "warnings": rebuild_combined(jid)["warnings"]}


# ── 작업 조회·조건 ────────────────────────────────────────────────────────────
@app.delete("/api/jobs/{jid}")
def delete_job(jid: str):
    if rt(jid)["busy"]:
        raise HTTPException(409, "진행 중인 작업은 삭제할 수 없습니다.")
    shutil.rmtree(jdir(jid))
    RUNTIME.pop(jid, None)
    return {"ok": True}


def list_files(o):
    out = []
    for dp, _, fs in os.walk(o):
        for f in sorted(fs):
            if not f.startswith("_"):
                out.append(os.path.relpath(os.path.join(dp, f), o).replace(os.sep, "/"))
    return sorted(out, key=lambda x: (x.count("/"), x))


@app.get("/api/jobs/{jid}")
def job(jid: str):
    m = meta(jid)
    if m.get("type") == "수량첨부":
        r = rt(jid)
        return {"id": jid, **{k: v for k, v in m.items() if k != "wiz"}, "files": list_files(outdir(jid)),
                "busy": r["busy"], "error": r["error"], "has_result": r["res"] is not None,
                "has_input": os.path.exists(cur_input(jid))}
    ov, ni, na = read_overview(cur_input(jid))
    r = rt(jid)
    subs = []
    cs = {c["name"]: c for c in m.get("combine_subs") or []}
    for i, s in enumerate(m.get("subs") or []):
        sov, sni, sna = read_overview(ipath(jid, s["file"]))
        subs.append({**s, "idx": i, "파일": os.path.basename(s["file"]), "파일_착공": sov.get("착공예정일"),
                     "적용_착공": cs.get(s["name"], {}).get("start") or sov.get("착공예정일"),
                     "공정수": sna, "수량수": sni, "총공사비": sov.get("총공사비(백만원)"), "시설유형": sov.get("시설유형(부록5)"),
                     "준비기간": sov.get("준비기간(일)"), "관측지점": sov.get("기상관측지점")})
    return {"id": jid, **m, "type": m.get("type", "단일공사"), "subs_info": subs, "composite_run": is_composite_run(m),
            "overview": {k: ov.get(k) for k in EDITABLE}, "n_items": ni, "n_acts": na,
            "files": list_files(outdir(jid)), "busy": r["busy"], "error": r["error"], "has_result": r["res"] is not None}


@app.post("/api/jobs/{jid}/select")
def select_input(jid: str, body: dict):
    m = meta(jid)
    if body.get("current") not in m["inputs"]:
        raise HTTPException(400, "입력 파일이 없습니다.")
    m["current"] = body["current"]
    save_meta(jid, m)
    rt(jid)["res"] = None
    return {"ok": True}


@app.post("/api/jobs/{jid}/settings")
def settings(jid: str, body: dict):
    m = meta(jid)
    body = {k: v for k, v in body.items() if k in EDITABLE}
    if is_composite_run(m) and m.get("kind") != "검토엑셀 변환":
        common = {k: v for k, v in body.items() if k not in PER_PROJECT}
        changed = []
        for s in m["subs"]:                       # 공통조건 → 모든 공사 입력서식
            changed += write_overview(ipath(jid, s["file"]), common)
        m["combined_overrides"] = {k: v for k, v in body.items() if k in ("공사명", "발주청", "사업위치") or k in common}
        if body.get("공사명"):
            m["name"] = body["공사명"]
        save_meta(jid, m)
        rebuild_combined(jid)
        changed = sorted(set(changed))
    else:
        changed = write_overview(cur_input(jid), body)
    rt(jid)["res"] = None
    return {"ok": True, "changed": changed}


# ── 실행 단계 ─────────────────────────────────────────────────────────────────
@app.post("/api/jobs/{jid}/run")
def run_review(jid: str):
    o = outdir(jid)

    def fn(r):
        r["log"] = RunLog(echo=False)
        r["res"] = None
        res = do_calc(jid, r["log"])
        r["log"]("[미리보기] 주공정 공정표 작성", "결과")
        service.chart(res, o)
        r["res"] = res
        s = service.summary(res)
        with open(os.path.join(o, "result.json"), "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False, indent=1, default=str)
        m = meta(jid)
        m["summary"] = {"공사기간_일": s["공사기간_일"], "공사기간_개월": s["공사기간_개월"], "적정성": s["적정성"]}
        save_meta(jid, m)
        r["log"](f"[완료] 공사기간 {s['공사기간_일']:,}일(≒{s['공사기간_개월']}개월), 적정성 '{s['적정성']}' — [결과 확인]에서 확인하십시오.", "완료")
    return background(jid, "검토 실행", fn)


@app.post("/api/jobs/{jid}/report")
def make_report(jid: str):
    o = outdir(jid)

    def fn(r):
        res = ensure_res(jid, r)
        r["log"]("[보고서 작성] 보고서·부록·공정표·검토과정 워크북" + (" + 공사별 보고서" if getattr(res, "subs", None) else ""), "보고서")
        service.reports(res, o)
        r["log"]("[완료] 보고서 작성 완료 — [보고서 출력]에서 PDF 변환·내려받기", "완료")
    return background(jid, "보고서 작성", fn)


@app.post("/api/jobs/{jid}/export")
def export(jid: str):
    o = outdir(jid)

    def fn(r):
        if not os.path.exists(os.path.join(o, "01_공사기간산정보고서.hwpx")):
            res = ensure_res(jid, r)
            r["log"]("보고서가 없어 먼저 작성합니다.", "출력")
            service.reports(res, o)
        r["log"] = r["log"] or RunLog(echo=False)
        r["log"]("[보고서 출력] PDF 변환 및 ZIP 묶음", "출력")
        service.export_pdf(o, r["log"])
        r["log"]("[완료] 보고서 출력 완료", "완료")
    return background(jid, "보고서 출력", fn)


@app.get("/api/jobs/{jid}/log")
def log(jid: str, since: int = 0):
    r = rt(jid)
    ent = r["log"].entries if r["log"] else []
    return {"entries": ent[since:], "next": len(ent), "busy": r["busy"], "error": r["error"]}


@app.get("/api/jobs/{jid}/result")
def result(jid: str):
    r = rt(jid)
    if r["res"] is not None:
        return JSONResponse(json.loads(json.dumps(service.summary(r["res"]), default=str)))
    p = os.path.join(outdir(jid), "result.json")
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return JSONResponse(json.load(f))
    raise HTTPException(404, "검토 결과가 없습니다. [검토 실행]을 먼저 하십시오.")


@app.get("/api/jobs/{jid}/file/{name:path}")
def file(jid: str, name: str, inline: int = 0):
    base = jdir(jid)
    for d in (outdir(jid), os.path.join(base, "입력"), os.path.join(base, "원본")):
        p = os.path.normpath(os.path.join(d, name))
        if p.startswith(os.path.normpath(d) + os.sep) and os.path.isfile(p):
            if inline:
                return FileResponse(p, headers={"Content-Disposition": "inline"},
                                    media_type="application/pdf" if p.endswith(".pdf") else None)
            return FileResponse(p, filename=os.path.basename(p))
    raise HTTPException(404, "파일이 없습니다.")


@app.get("/api/jobs/{jid}/hwpx_html/{name:path}", response_class=HTMLResponse)
def hwpx_html(jid: str, name: str):
    """HWPX 보고서 화면 미리보기(텍스트·표 구성 확인용 HTML)."""
    p = os.path.normpath(os.path.join(outdir(jid), name))
    if not p.startswith(os.path.normpath(outdir(jid)) + os.sep) or not p.endswith(".hwpx") or not os.path.isfile(p):
        raise HTTPException(404, "HWPX 파일이 없습니다.")
    from hwpx import HwpxDocument
    doc = HwpxDocument.open(p)
    try:
        html = doc.export_html()
    finally:
        doc.close()
    css = ("<style>body{font-family:'Malgun Gothic','맑은 고딕','Noto Sans KR',sans-serif;max-width:900px;margin:20px auto;"
           "padding:0 20px;font-size:14px;line-height:1.6}table{border-collapse:collapse;margin:6px auto;font-size:12.5px}"
           "td{border:1px solid #999;padding:3px 6px}tr:first-child td{background:#d9e2f3;font-weight:700;text-align:center}"
           "hr{border:0;border-top:3px double #1f4e78;margin:30px 0}</style>")
    return HTMLResponse(html.replace("</head>", css + "</head>", 1))


wiz.CTX.update(jdir=jdir, meta=meta, save_meta=save_meta, new_job_dir=new_job_dir, rt=rt)
