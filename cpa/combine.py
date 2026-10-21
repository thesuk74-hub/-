"""복합공사 — 공사별 입력서식(수량·공정)을 하나의 통합 입력서식으로 합친다.

 · 공정ID  : <공사코드>-<원래ID> (예: A-C10), 공정명 앞에 [공사명]
 · 선행공정: 같은 공사 안의 ID를 자동 치환
 · 착공 시차: 후속 공사의 선행공정이 없는 공정에 '착수가능일'(= 공사 착공일 [+ 별도 준비기간]) 지정
 · 준비기간: 최초 착공 공사의 준비기간 1회, 정리기간: 공사별 정리기간 중 최대값 1회 (검토엑셀 ⑥ 방식)
 · 비작업조건(2_비작업조건): 최초 착공 공사 기준 (공사별 기상프로필이 다르면 경고)
공사별 적정성(부록5·동종실적)은 각 공사의 입력서식으로 독립 산정한다(검토엑셀 ⑦ 방식).
"""
import datetime as dt
import re
import string

import openpyxl

from .project import ACT_COLS, ITEM_COLS, read_project, write_template

CODES = string.ascii_uppercase


def _sheet_rows(ws, cols):
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {k: (row[idx[k]] if k in idx and idx[k] < len(row) else None) for k in cols}
        if any(v not in (None, "") for v in d.values()):
            out.append(d)
    return out


def _overview_raw(wb):
    ov = {}
    for row in wb["1_공사개요"].iter_rows(min_row=2, values_only=True):
        if row[0] and not str(row[0]).startswith("■"):
            ov[str(row[0]).strip()] = row[1]
    return ov


def _profiles_raw(wb):
    ws = wb["2_비작업조건"]
    pn = [h for h in [c.value for c in ws[2]][3:] if h]
    prof, extra = {p: {} for p in pn}, {}
    hoff, wex = set(), set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if isinstance(row[0], (int, float)) and 1 <= row[0] <= 15:
            for j, p in enumerate(pn):
                if row[3 + j] not in (None, ""):
                    prof[p][int(row[0])] = float(row[3 + j])
        if row[2] == "법정공휴일 적용":
            hoff = {p for j, p in enumerate(pn) if str(row[3 + j] or "Y").upper().startswith("N")}
        if row[2] == "동절기 중지 적용":
            wex = {p for j, p in enumerate(pn) if str(row[3 + j] or "Y").upper().startswith("N")}
        if row[2] in ("발령일수", "사용자정의"):
            extra[row[2]] = ([float(x or 0) for x in row[3:15]], row[1])
    return prof, extra, hoff, wex


def build(specs, out_path, name):
    """specs: [{"name": 공사명, "path": 입력서식, "start": date|None, "own_prep": bool}] → 통합 입력서식 작성.

    반환: {"warnings": [...], "subs": [{"code", "name", "start", "snet", "acts", "items", "cost"}]}
    """
    warns, subs = [], []
    loaded = []
    for k, s in enumerate(specs):
        wb = openpyxl.load_workbook(s["path"], data_only=True)
        p = read_project(s["path"])          # 형식 검증
        ov = _overview_raw(wb)
        start = s.get("start") or p.start_date
        if isinstance(start, str):
            start = dt.date.fromisoformat(start[:10])
        loaded.append(dict(spec=s, wb=wb, p=p, ov=ov, start=start, code=CODES[k]))
    loaded.sort(key=lambda x: x["start"])
    base = loaded[0]
    prof, extra, hoff, wex = _profiles_raw(base["wb"])
    for L in loaded[1:]:
        pr2, ex2, _, _ = _profiles_raw(L["wb"])
        for pn, c in pr2.items():
            if pn not in prof:
                prof[pn] = c
                warns.append(f"[{L['spec']['name']}] 기상프로필 '{pn}' 추가")
            elif prof[pn] != c:
                warns.append(f"[{L['spec']['name']}] 기상프로필 '{pn}' 조건이 최초 공사와 달라 최초 공사 기준 적용")
        if ex2.get("사용자정의", ([0] * 12,))[0] != extra.get("사용자정의", ([0] * 12,))[0]:
            warns.append(f"[{L['spec']['name']}] 사용자정의 기상조건이 최초 공사와 달라 최초 공사 기준 적용")
        if L["p"].g("기상관측지점") != base["p"].g("기상관측지점"):
            warns.append(f"[{L['spec']['name']}] 관측지점 '{L['p'].g('기상관측지점')}' ≠ 최초 공사 '{base['p'].g('기상관측지점')}' — 통합 산정은 최초 공사 관측지점 적용")

    items, acts = [], []
    for L in loaded:
        code, nm, wb = L["code"], L["spec"]["name"], L["wb"]
        tag = f"[{nm}]"
        a_rows = _sheet_rows(wb["4_공정"], ACT_COLS)
        ids = {str(a["공정ID"]).strip() for a in a_rows if a["공정ID"]}
        own_prep = bool(L["spec"].get("own_prep"))
        snet = None
        if L is not base:
            snet = L["start"] + dt.timedelta(days=L["p"].prep_days if own_prep else 0)

        def rid(x):
            return f"{code}-{x}"

        for a in a_rows:
            if not a["공정ID"]:
                continue
            a = dict(a)
            old = str(a["공정ID"]).strip()
            a["공정ID"] = rid(old)
            nm0 = str(a["공정명"] or "")
            a["공정명"] = nm0 if nm0.startswith("[") else f"{tag} {nm0}".strip()
            a["소속공사"] = nm
            preds = []
            for tok in re.split(r"[,;\n]", str(a["선행공정"] or "")):
                tok = tok.strip()
                if not tok:
                    continue
                pid, _, rest = tok.partition(":")
                if pid not in ids:
                    warns.append(f"[{nm}] 공정 {old}의 선행공정 '{pid}' 없음 — 제외")
                    continue
                preds.append(rid(pid) + (":" + rest if rest else ""))
            a["선행공정"] = ", ".join(preds)
            if snet and not preds:
                a["착수가능일"] = max(snet, a["착수가능일"].date() if isinstance(a.get("착수가능일"), dt.datetime) else (a.get("착수가능일") or snet))
            acts.append(a)
        n_items = 0
        for it in _sheet_rows(wb["3_수량"], ITEM_COLS):
            if not it["공정ID"]:
                continue
            it = dict(it)
            it["공정ID"] = rid(str(it["공정ID"]).strip())
            it["No"] = len(items) + 1
            it["비고"] = f"{tag} " + str(it["비고"] or "")
            items.append(it)
            n_items += 1
        subs.append(dict(code=code, name=nm, start=str(L["start"]), snet=str(snet) if snet else None,
                         acts=len(a_rows), items=n_items, cost=L["p"].total_cost, path=L["spec"]["path"]))

    ov = dict(base["ov"])
    ov.update({"공사명": name, "총공사비(백만원)": sum(L["p"].total_cost for L in loaded), "시설유형(부록5)": "해당없음",
               "착공예정일": base["start"], "준비기간(일)": base["p"].prep_days,
               "정리기간(일)": max(L["p"].cleanup_days for L in loaded),
               "공사특성": "복합공사 — " + " / ".join(f"{s['name']}(착공 {s['start']}, 공정 {s['acts']}개·수량 {s['items']}항목)" for s in subs)
                       + "\n준비기간은 최초 착공 공사에, 정리기간은 최종 준공 시점에 1회 계상"
                       + "\n후속 공사는 착공일(착수가능일) 이후 착수, 공사별 적정성은 공사별 입력으로 독립 검토"})
    for k in ("연면적(100㎡)", "지상층수", "지하층수", "하수도연장(m)", "도로연장(m)", "관경(mm)", "시설용량(㎥/일)"):
        ov.pop(k, None)
    ov["환경공단 산정식"] = "해당없음"
    ov["환경공단 경험치"] = "해당없음"
    dust = extra.get("발령일수", ([0] * 12, None))[0]
    custom, clabel = extra.get("사용자정의", ([0] * 12, "사용자정의"))
    write_template(out_path, ov, prof, dust, items, acts, [], custom=custom, custom_label=clabel or "사용자정의",
                   holiday_off=hoff, winter_exempt=wex)
    return {"warnings": warns, "subs": subs}
