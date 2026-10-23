"""엑셀 부록 작성 도우미 — docbuild.Doc 과 같은 인터페이스. 부록(장·구역)별 시트, 표는 셀 표, 그림은 이미지."""
import re

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NAVY = "1F4E78"


class XlsxDoc:
    def __init__(self, project_name):
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "표지"
        self.pname = project_name
        self.r = 1
        self.ch = self.tn = self.fn = 0
        self.avail = 26.7
        self._widths = {}
        self.index = []

    # ── 시트 ──
    def _sheet(self, title):
        self._flush_widths()
        base = re.sub(r"[\\/*?:\[\]]", "", title)[:28] or "시트"
        name, k = base, 2
        while name in self.wb.sheetnames:
            name = f"{base[:26]}_{k}"
            k += 1
        self.ws = self.wb.create_sheet(name)
        self.ws.sheet_view.showGridLines = False
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.paperSize = self.ws.PAPERSIZE_A4
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_setup.fitToWidth, self.ws.page_setup.fitToHeight = 1, 0
        self.r = 1
        self._widths = {}
        c = self.ws.cell(1, 1, f"{self.pname}  |  {title}")
        c.font = Font(size=9, color="555555", italic=True)
        self.r = 3
        self.index.append((name, title))
        return self.ws

    def _flush_widths(self):
        for j, w in self._widths.items():
            self.ws.column_dimensions[get_column_letter(j)].width = max(6, min(60, w))

    def _w(self, j, chars):
        self._widths[j] = max(self._widths.get(j, 8), chars)

    def first_header(self, right_text):
        self.ws.cell(1, 1, f"{self.pname}  |  {right_text}").font = Font(size=9, color="555555", italic=True)

    def set_header(self, section, right_text):
        pass

    def new_section(self, header_text, landscape=False):
        return self._sheet(header_text)

    # ── 본문 ──
    def _line(self, text, size=10, bold=False, color=None, indent=0):
        c = self.ws.cell(self.r, 1 + (1 if indent and indent >= 1.0 else 0), text)
        c.font = Font(size=size, bold=bold, color=color)
        c.alignment = Alignment(vertical="top")
        self.r += 1
        return c

    def cover(self, title, sub, date_text, org, extra=None):
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 100
        for i, (t, sz, b) in enumerate(((title, 18, True), (sub, 24, True), (extra or "", 11, False), ("", 10, False),
                                        (date_text, 14, True), (org, 14, True)), 4):
            c = ws.cell(i * 2, 1, t)
            c.font = Font(size=sz, bold=b, color=NAVY if sz == 24 else None)
            c.alignment = Alignment(horizontal="center")

    def toc(self, entries):
        self._line("< 목  차 >", 14, True)
        for lvl, text in entries:
            self._line(("    " * lvl) + text, 11 if lvl == 0 else 10, lvl == 0)
        self.r += 1

    def chapter(self, no, title, subs=(), title_page=True):
        self.ch, self.tn, self.fn = no, 0, 0
        self._sheet(f"{no}장 {title}" if isinstance(no, int) else f"{no} {title}")
        self._line(f"제 {no} 장  {title}", 16, True, NAVY)
        for s in subs:
            self._line("  " + s, 11)
        self.r += 1

    def h2(self, text):
        self.r += 1
        self._line(text, 13, True, NAVY)

    def h3(self, text):
        self._line(text, 11.5, True)

    def h4(self, text):
        self._line(text, 10.5, True)

    def para(self, text, indent=1.2, size=10, bold=False, italic=False, color=None):
        col = None if color is None else str(color).replace("#", "")[-6:]
        for ln in str(text).split("\n"):
            self._line("  " + ln, size, bold, col)

    def bullet(self, text, indent=1.2, mark="-"):
        self._line(f"   {mark} {text}", 10)

    def note(self, text, indent=0.8):
        self._line("  " + text, 9, False, "444444")

    def boxed(self, lines, title=None, fill="F2F2F2"):
        if title:
            self._line(title, 10, True, NAVY)
        for ln in lines:
            c = self._line(ln, 10.5 if len(lines) == 1 else 9.5, len(lines) == 1)
            c.fill = PatternFill("solid", fgColor=fill)
        self.r += 1

    def caption_table(self, title):
        self.tn += 1
        self._line(f"[표 {self.ch}-{self.tn}]  {title}", 10, True)

    def table(self, header, rows, title=None, widths=None, size=8.5, align=None, hdr_fill="D9E2F3", bold_rows=(),
              fills=None, header_rows=None):
        if title:
            self.caption_table(title)
        hdrs = header_rows or [header]
        ncol = len(hdrs[0])
        r0 = self.r
        for hr in hdrs:
            for j, h in enumerate(hr, 1):
                c = self.ws.cell(self.r, j, h)
                c.font, c.border = Font(size=size + 0.5, bold=True), BOX
                c.fill = PatternFill("solid", fgColor=hdr_fill)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            self.r += 1
        for i, row in enumerate(rows):
            for j, v in enumerate(row[:ncol], 1):
                c = self.ws.cell(self.r, j, v)
                c.font, c.border = Font(size=size + 0.5, bold=i in bold_rows), BOX
                a = align[j - 1] if align else ("L" if isinstance(v, str) and len(v) > 14 else "C")
                c.alignment = Alignment(horizontal={"L": "left", "C": "center", "R": "right"}[a], vertical="center",
                                        wrap_text=isinstance(v, str) and len(v) > 40)
                if isinstance(v, float):
                    c.number_format = "#,##0.0" if abs(v - round(v)) > 1e-9 else "#,##0"
                elif isinstance(v, int):
                    c.number_format = "#,##0"
                if fills and fills[i]:
                    c.fill = PatternFill("solid", fgColor=fills[i])
            self.r += 1
        for j in range(ncol):
            if widths and j < len(widths):
                self._w(j + 1, widths[j] * 4.2)
            else:
                ln = max([len(str(hr[j])) for hr in hdrs if j < len(hr)] + [len(str(rw[j])) for rw in rows if j < len(rw)] + [4])
                self._w(j + 1, min(ln * 1.3 + 2, 45))
        if len(hdrs) == 1:
            self.ws.auto_filter.ref = None
        self.r += 1
        return (r0, self.r)

    def figure(self, path, title, width_cm=16, max_h_cm=None):
        from PIL import Image
        w, h = Image.open(path).size
        if max_h_cm:
            width_cm = min(width_cm, max_h_cm * w / h)
        img = XLImage(path)
        px_w = width_cm / 2.54 * 96
        img.width, img.height = px_w, px_w * h / w
        self.ws.add_image(img, f"A{self.r}")
        self.r += int(img.height / 20) + 2
        self.fn += 1
        self._line(f"[그림 {self.ch}-{self.fn}] {title}", 10, True)

    def page_break(self):
        self.r += 1

    def save(self, path):
        self._flush_widths()
        if self.index:
            ws = self.wb["표지"]
            r = max(ws.max_row + 2, 18)
            ws.cell(r, 1, "시트 목록").font = Font(bold=True, size=11)
            for i, (name, title) in enumerate(self.index, 1):
                c = ws.cell(r + i, 1, f"{i}. {title}")
                c.hyperlink = f"#'{name}'!A1"
                c.font = Font(color="0563C1", underline="single")
        self.wb.save(path)
